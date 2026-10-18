from pathlib import Path
from tempfile import TemporaryDirectory
from pywpdoc import Document
import docx  # confirm optional python-docx dependency

with TemporaryDirectory() as tmp:
    for fmt in ("docx", "hwpx"):
        p = Path(tmp) / f"smoke.{fmt}"
        d = Document.new(fmt)
        d.add_paragraph("wheel installation smoke")
        d.add_table(1, 1).cell(0, 0).text = "safe"
        d.save(p)
        reopened = Document.open(p)
        assert "wheel installation smoke" in [par.text for par in reopened.paragraphs]
        assert reopened.tables[0].cell(0, 0).text == "safe"
print("Wheel smoke passed: DOCX + HWPX")
