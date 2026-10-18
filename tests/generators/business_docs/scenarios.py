"""Deterministic, synthetic, redistributable business-document scenarios.

IMPORTANT: Only PyWPDoc's public Common Document API creates the documents.
No source fixtures, private sample documents, or backend-native authoring APIs.
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import struct
import zlib

from pywpdoc import Document
from pywpdoc._engine import Pt, WD_STYLE_TYPE


@dataclass(frozen=True)
class Scenario:
    key: str
    title: str
    domain: str
    sections: tuple[str, ...]
    columns: tuple[str, ...]
    records: tuple[tuple[str, ...], ...]
    image: bool = False
    nested: bool = False
    merged: bool = False
    template: bool = False
    long: bool = False
    extra_tables: int = 0
    appendix_pages: int = 0


def _s(key, title, domain, sections, columns, *records, **flags):
    return Scenario(key, title, domain, tuple(sections.split('|')), tuple(columns.split('|')),
                    tuple(tuple(row.split('|')) for row in records), **flags)


SCENARIOS: tuple[Scenario, ...] = (
    _s('01_contract', '용역 계약서', '법무', '계약 목적|업무 범위|계약 기간|검수 조건', '항목|약정', '기간|2026년 4분기', '금액|12,000,000원', '검수|양측 서명', merged=True),
    _s('02_nda', '비밀유지계약서', '법무', '비밀정보의 정의|수령자의 의무|예외 조항|반환 절차', '구분|보호 의무', '소스 코드|외부 제공 금지', '고객 자료|암호화 보관'),
    _s('03_mou', '업무협약 양해각서', '법무', '협력 배경|상호 역할|성과 공유|유효기간', '기관|역할', '갑 기관|연구 자료 제공', '을 기관|품질 검증'),
    _s('04_official_notice', '협조 요청 공문', '행정', '수신 및 참조|요청 사항|기한|첨부 안내', '번호|안내 내용', '1|2026-10-31까지 회신', '2|담당 부서 확인'),
    _s('05_meeting_minutes', '주간 회의록', '운영', '참석자|주요 안건|결정 사항|후속 조치', '안건|담당|기한', '품질 게이트|QA팀|금요일', '출시 승인|PM|월요일'),
    _s('06_project_kickoff', '프로젝트 착수 보고서', '프로젝트', '목표|범위|담당자|일정', '단계|시작일|종료일', '설계|2026-10-01|2026-10-15', '구현|2026-10-16|2026-11-15'),
    _s('07_monthly_report', '월간 업무 실적 보고서', '운영', '목표 대비 실적|이슈|개선 계획|다음 달 목표', '지표|목표|실적', '처리 건수|120|128', '만족도|90%|93%', image=True),
    _s('08_quarterly_review', '분기 성과 검토', '경영', '개요|성과 분석|위험 요인|차기 전략', '사업|전분기|금분기', 'A|100|127', 'B|80|85', image=True),
    _s('09_expense_report', '출장 경비 정산서', '회계', '출장 목적|일정|비용 내역|승인', '지출|금액|증빙', '교통비|82,000원|영수증 01', '식비|36,000원|영수증 02'),
    _s('10_invoice', '세금계산서 유사 청구 내역', '회계', '청구 대상|공급 내역|합계|지급 안내', '품목|수량|단가', '서비스 A|2|300,000원', '서비스 B|1|500,000원'),
    _s('11_purchase_order', '구매 발주서', '구매', '구매 목적|발주 품목|납기|인수 검사', '품목|단위|수량', '노트북|대|8', '모니터|대|16', merged=True),
    _s('12_procurement_request', '구매 품의서', '구매', '요청 사유|수량 및 단가|공급사 후보|결재', '후보|단가|납기', '업체 A|150,000원|7일', '업체 B|142,000원|14일'),
    _s('13_onboarding', '신규 입사자 안내서', '인사', '입사 첫날|계정 신청|보안 교육|담당자', '절차|확인', '사원증 발급|완료', '보안 서약|예정', template=True),
    _s('14_hr_policy', '휴가 및 재택근무 규정', '인사', '적용 대상|승인 절차|예외 사항|개정 이력', '분류|기준', '연차|사전 신청', '재택|팀장 승인'),
    _s('15_incident_report', '서비스 장애 보고서', 'IT', '발생 시각|영향 범위|원인|재발 방지', '시각|이벤트|대응', '09:12|접속 오류 감지|당직자 호출', '09:45|복구 완료|모니터링'),
    _s('16_release_notes', '소프트웨어 릴리스 노트', '개발', '주요 변경|버그 수정|호환성|롤백 방법', 'ID|변경 내용', 'BUG-101|표 속성 검사', 'BUG-102|재구축 보호'),
    _s('17_user_manual', '업무 시스템 사용자 매뉴얼', '기술문서', '설치|시작하기|단계별 사용법|문제 해결', '단계|작업', '1|로그인', '2|프로젝트 생성', image=True, nested=True, appendix_pages=3),
    _s('18_sop', '표준 작업 절차서', '품질', '목적|준비물|안전 주의사항|작업 순서', '순서|작업|점검', '1|장비 전원 확인|PASS', '2|시료 채취|PASS', nested=True),
    _s('19_checklist', '운영 점검표', '운영', '점검 범위|일일 점검|월간 점검|후속 조치', '항목|결과|담당', '백업|완료|운영팀', '로그 보존|완료|보안팀'),
    _s('20_audit_report', '감사 결과 보고서', '감사', '감사 기준|주요 발견|개선 권고|후속 점검', '위험|심각도|담당', '접근 제어|높음|보안팀', '로그 보존|중간|운영팀', appendix_pages=2),
    _s('21_lab_report', '시험·분석 보고서', '연구', '목적|시험 방법|측정값|결론', '시료|온도|결과', '시료 A|23.5℃|적합', '시료 B|24.1℃|적합', image=True),
    _s('22_school_notice', '학교 가정통신문', '교육', '학부모 안내|행사 일정|준비물|문의', '일정|내용', '2026-11-02|체험 학습', '2026-11-03|소감 제출', template=True),
    _s('23_complex_tables', '재고·자산 관리 대장', '자산', '관리 범위|등록 자산|위치 변경|재물조사', '자산 번호|품목|수량|위치', 'A-001|노트북|20|본사', 'A-002|모니터|35|지사', 'A-003|스위치|8|서버실', nested=True, merged=True, extra_tables=2),
    _s('24_long_form', '대용량 종합 검토 보고서', '품질', '검토 목적|계층별 점검|장문 문단|요약', '검사항목|결과|설명', '무결성|PASS|구조 검증', '호환성|PASS|재열기 검증', long=True, image=True, extra_tables=12),
)


def test_image_bytes(color: tuple[int, int, int] = (20, 92, 161), width: int = 220, height: int = 66) -> bytes:
    """Portable PNG test graphic; no external image assets required."""
    def chunk(tag: bytes, raw: bytes) -> bytes:
        return struct.pack('>I', len(raw)) + tag + raw + struct.pack('>I', zlib.crc32(tag + raw) & 0xffffffff)
    scan = b''.join(b'\x00' + bytes(color) * width for _ in range(height))
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>2I5B', width, height, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(scan, 6)) + chunk(b'IEND', b'')


def _set_comparable_qa_paragraph_spacing(doc, fmt: str) -> None:
    """Make synthetic QA samples comparable without changing product defaults.

    Word's Normal default carries 10 pt paragraph-after spacing and 11 pt text;
    the HWPX baseline has 0 pt paragraph-after spacing and 10 pt text.  The
    resulting page-count difference is a *fixture authoring* discrepancy, not
    evidence of a conversion regression.  Use 10 pt / 16 pt line advance / 0 pt
    paragraph spacing as the explicit comparison baseline.  In Hancom this
    corresponds to 160% (of 10 pt).  For Word use an exact physical advance,
    because a nominal Word 1.6 multiple measures larger than 16 pt with the
    tested fonts.  Do not change the backends' default document semantics.

    The DOCX backend currently uses python-docx paragraph formatting; import
    only its public Pt value type for exact physical units, not its document
    constructor or any low-level XML authoring API.  Native visual approval is
    still required on Windows.
    """
    normal = doc.styles.default(WD_STYLE_TYPE.PARAGRAPH)
    if normal is None:
        raise RuntimeError("missing default paragraph style")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    if fmt == "docx":
        from docx.shared import Pt as DocxPt
    for paragraph in doc.paragraphs:
        if fmt == "docx":
            paragraph.paragraph_format.line_spacing = DocxPt(16)
        elif fmt == "hwpx":
            paragraph.paragraph_format.line_spacing_percent = 160
        else:
            raise ValueError(f"unsupported QA format: {fmt!r}")


def _set_comparable_qa_page_geometry(doc, fmt: str) -> None:
    """Normalize *effective* page geometry for synthetic cross-format QA.

    Hancom reserves header and footer distances in addition to the top/bottom
    page margins, even when both stories are empty.  Word's header/footer
    distances do not shrink its main body that way.  Both backends' raw A4
    width and all four margins already match.  For headerless/footerless QA
    documents only, remove HWPX's unused reserved header/footer areas.

    Do not alter public product defaults or documents with actual headers.
    The HWPX section-anchor paragraph can still offset first-page content;
    native layout comparison is required before accepting visual parity.
    """
    if fmt == 'hwpx':
        doc.page_setup.header_distance = 0
        doc.page_setup.footer_distance = 0
    elif fmt != 'docx':
        raise ValueError(f"unsupported QA format: {fmt!r}")


def build_document(s: Scenario, fmt: str):
    """Author via the Common API only; return document + expected observable contents."""
    doc = Document.new(fmt)
    _set_comparable_qa_page_geometry(doc, fmt)
    doc.core_properties.title = s.title
    doc.core_properties.author = 'PyWPDoc synthetic QA'
    doc.core_properties.subject = s.domain
    sentinel = f'QA-SCENARIO::{s.key}::{fmt}::END'
    paragraph_expect: list[str] = []
    table_expect: list[tuple[int, int, str]] = []

    def para(value: str, *, bold: bool = False):
        p = doc.add_paragraph()
        r = p.add_run(value)
        if bold:
            r.bold = True
        paragraph_expect.append(value)
        return p

    para(s.title, bold=True)
    para(f'문서번호 {s.key.upper()} · 작성일 2026-10-15 · 부서 {s.domain}')
    for i, heading in enumerate(s.sections, 1):
        para(f'{i}. {heading}', bold=True)
        para(f'{heading}에 대한 검토 및 조치 사항을 기록합니다. 담당 부서는 {s.domain}이며, 근거 문서는 내부 QA 합성 자료입니다.')

    table = doc.add_table(len(s.records) + 1, len(s.columns), style="Table Grid")
    for c, name in enumerate(s.columns):
        table.cell(0, c).text = name
        table_expect.append((0, c, name))
    for r, row in enumerate(s.records, 1):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
            table_expect.append((r, c, value))

    if s.nested:
        nested = table.cell(1, 0).add_table(2, 2, style="Table Grid")
        nested.cell(0, 0).text = '중첩-1'
        nested.cell(0, 1).text = 'Nested-A'
        nested.cell(1, 0).text = '중첩-2'
        nested.cell(1, 1).text = 'Nested-B'
    if s.merged:
        # Use an extra independent table so original data cells retain stable coordinates.
        # Separate independent tables with a real top-level paragraph.
        # Word may coalesce adjacent w:tbl blocks on native re-save.
        if fmt == 'docx':
            doc.add_paragraph('')
        band = doc.add_table(1, 2, style="Table Grid")
        band.cell(0, 0).text = '병합-확인-MERGE-MARK'
        band.cell(0, 0).merge(band.cell(0, 1))
    if s.template:
        p = doc.add_paragraph('담당자 {{person}} · 회신기한 {{deadline}}')
        doc.render_template({'person': '홍길동', 'deadline': '2026-11-30'}, everywhere=False, missing='error')
        paragraph_expect.append('담당자 홍길동 · 회신기한 2026-11-30')
    if s.image:
        doc.add_picture(test_image_bytes(), 'png')
        para('이미지 첨부: QA-IMAGE-MARK (합성 PNG, 실제 업무 사진 아님)')
    for n in range(s.extra_tables):
        # Even an empty w:p distinguishes independent OOXML tables.
        # Do not rely on API-level table counts only.
        if fmt == 'docx':
            doc.add_paragraph('')
        extra = doc.add_table(2, 2, style="Table Grid")
        extra.cell(0, 0).text = f'EXTRA-{n:02d}'
        extra.cell(1, 1).text = f'DATA-{n:02d}'
    if s.appendix_pages:
        for page in range(s.appendix_pages):
            doc.add_page_break()
            para(f"부록 {page+1}: {s.title} 점검 절차와 참고 자료", bold=True)
            for item in range(16):
                para(f"부록-{page+1:02d}-{item+1:02d}: 검토 근거 · 실제 시스템과 입력 서식을 확인합니다.")
    if s.long:
        for n in range(2750):
            para(f'장문-{n:04d}: The quick brown fox · 문서 검증 · العربية · 日本語 · 데이터 보존.')
    para(sentinel)
    _set_comparable_qa_paragraph_spacing(doc, fmt)
    return doc, {'sentinel': sentinel, 'paragraphs': paragraph_expect,
                 'first_table_cells': table_expect, 'has_image': s.image,
                 'has_nested': s.nested, 'has_merge': s.merged, 'extra_tables': s.extra_tables,
                 'tables': 1 + int(s.merged) + s.extra_tables,
                 'id': s.key, 'format': fmt, 'title': s.title}


def open_and_check(path: Path, expected: dict) -> dict:
    doc = Document.open(path)
    assert doc.format == expected['format']
    assert doc.core_properties.title == expected['title']
    got = [p.text for p in doc.paragraphs]
    # Native image placement may include an additional empty paragraph.
    for marker in expected['paragraphs']:
        assert marker in got, f'missing paragraph {marker[:60]} in {path}'
    assert got.count(expected['sentinel']) == 1
    assert len(doc.tables) == expected['tables']
    if expected['format'] == 'docx' and expected['tables'] > 1:
        from tests.generators.business_docs.structure import audit_docx_independent_tables
        audit_docx_independent_tables(path, expected_tables=expected['tables'])
    t = doc.tables[0]
    for row, col, value in expected['first_table_cells']:
        assert value in t.cell(row, col).text, (row, col, value, t.cell(row, col).text)
    if expected['has_nested']:
        nested = t.cell(1, 0).tables
        assert len(nested) == 1
        assert nested[0].cell(1, 1).text == 'Nested-B'
    if expected['has_merge']:
        assert 'MERGE-MARK' in doc.tables[1].cell(0, 0).text
    first_extra = 1 + int(expected['has_merge'])
    for index in range(expected.get('extra_tables', 0)):
        extra = doc.tables[first_extra + index]
        assert extra.cell(0, 0).text == f'EXTRA-{index:02d}'
        assert extra.cell(1, 1).text == f'DATA-{index:02d}'
    return {'paragraphs': len(got), 'tables': len(doc.tables), 'first_table_cells': len(expected['first_table_cells'])}
