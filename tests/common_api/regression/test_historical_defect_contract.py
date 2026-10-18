"""Behavior-oriented guards for defects reported in independent reviews.

The larger existing suites exercise backend-specific XML. These ensure the
public API retains the fixed semantics on future implementations.
"""
from __future__ import annotations
import zipfile
import pytest
from pywpdoc import Document
from pywpdoc._engine.core.errors import DocumentUnsupportedError, DocumentValueError


@pytest.mark.common_api
@pytest.mark.regression
@pytest.mark.parametrize("body", ["고객 정보", "매출 2026", "Policy section"])
def test_issue_docx_rebuild_never_silently_empties_data(tmp_path, body):
    doc = Document.new("docx")
    doc.add_paragraph(body)
    doc.add_table(2, 2).cell(1, 1).text = "KEEP CELL"
    src = tmp_path / "before.docx"
    dst = tmp_path / "after.docx"
    doc.save(src)
    Document.rebuild(src, format="docx").save(dst)
    with zipfile.ZipFile(src) as first, zipfile.ZipFile(dst) as second:
        assert first.read("word/document.xml") == second.read("word/document.xml")
    reopened = Document.open(dst)
    assert reopened.paragraphs[0].text == body
    assert reopened.tables[0].cell(1, 1).text == "KEEP CELL"


@pytest.mark.common_api
@pytest.mark.regression
def test_issue_docx_to_hwpx_must_fail_explicitly(tmp_path):
    doc = Document.new("docx")
    doc.add_paragraph("important")
    src = tmp_path / "in.docx"
    doc.save(src)
    with pytest.raises(DocumentUnsupportedError) as exc:
        Document.rebuild(src, format="hwpx")
    assert exc.value.code == "rebuild-docx-to-hwpx-unsupported"
    assert Document.open(src).paragraphs[0].text == "important"


@pytest.mark.common_api
@pytest.mark.regression
@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
def test_issue_compare_ignores_nothing_outside_requested_scope(fmt):
    a, b = Document.new(fmt), Document.new(fmt)
    a.add_paragraph("same")
    b.add_paragraph("same")
    a.add_table(1, 1).cell(0, 0).text = "OLD"
    b.add_table(1, 1).cell(0, 0).text = "NEW"
    with pytest.raises(DocumentUnsupportedError) as exc:
        a.compare(b, everywhere=False)
    assert exc.value.code == "compare-outside-scope-changes-unsupported"
    assert a.tables[0].cell(0, 0).text == "OLD"
    assert b.tables[0].cell(0, 0).text == "NEW"


@pytest.mark.common_api
@pytest.mark.regression
@pytest.mark.parametrize("border_scope", ["table", "cell"])
def test_issue_docx_border_only_outside_scope_fails_closed(border_scope):
    a, b = Document.new("docx"), Document.new("docx")
    for doc in (a, b):
        doc.add_paragraph("same")
        doc.add_table(1, 1).cell(0, 0).text = "same"
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    if border_scope == "table":
        pr = b.docx.tables[0]._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        borders.append(bottom)
        pr.append(borders)
    else:
        tcpr = b.docx.tables[0].cell(0, 0)._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        borders.append(bottom)
        tcpr.append(borders)
    with pytest.raises(DocumentUnsupportedError) as exc:
        a.compare(b, everywhere=False)
    assert exc.value.code == "compare-outside-scope-changes-unsupported"


@pytest.mark.common_api
@pytest.mark.regression
@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
def test_issue_template_preflight_avoids_partial_document_mutation(fmt):
    doc = Document.new(fmt)
    doc.add_paragraph("First {{ok}}")
    doc.add_paragraph("Second {{missing}}")
    with pytest.raises(DocumentValueError) as exc:
        doc.render_template({"ok": "OK"}, missing="error", everywhere=False)
    assert exc.value.code == "template-value-missing"
    assert [p.text for p in doc.paragraphs] == ["First {{ok}}", "Second {{missing}}"]


@pytest.mark.common_api
@pytest.mark.regression
@pytest.mark.parametrize("text", ["검증 셀", "밑 행 텍스트"])
def test_issue_hwpx_to_docx_table_not_forced_to_exact_tiny_height(text, tmp_path):
    from lxml import etree
    doc = Document.new("hwpx")
    doc.add_paragraph("본문")
    doc.add_table(2, 1).cell(0, 0).text = text
    out = tmp_path / "render.docx"
    Document.rebuild(doc, format="docx").save(out)
    with zipfile.ZipFile(out) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    rows = root.xpath("//w:tr[.//w:t[text()=$value]]", namespaces=ns, value=text)
    assert rows
    assert not rows[0].xpath("./w:trPr/w:trHeight[@w:hRule='exact']", namespaces=ns)


@pytest.mark.common_api
@pytest.mark.regression
def test_issue_hwpx_large_batch_replacement_does_not_skip_late_items():
    doc = Document.new("hwpx")
    for i in range(500):
        doc.add_paragraph(f"key{i:04d} old")
    assert doc.replace_text("old", "new") == 500
    assert doc.paragraphs[-1].text.endswith("new")
