"""Behavioral parity oracles from the *public* python-docx API (not copied tests).

This tests a limited feature intersection only; it is NOT full python-docx parity.
"""
from __future__ import annotations
import pytest
from docx import Document as ReferenceDocument
from pywpdoc import Document


def _save_reopen_both(reference, candidate, fmt, tmp_path):
    r = tmp_path / "reference.docx"
    c = tmp_path / ("candidate." + fmt)
    reference.save(r)
    candidate.save(c)
    return ReferenceDocument(r), Document.open(c)


@pytest.mark.common_api
@pytest.mark.parity
@pytest.mark.parametrize("texts", [("A", "B", "C"), ("한국어", "English", "中文"), ("", "one")])
@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
def test_public_paragraph_sequence_matches_python_docx(fmt, texts, tmp_path):
    ref, cand = ReferenceDocument(), Document.new(fmt)
    for text in texts:
        ref.add_paragraph(text)
        cand.add_paragraph(text)
    reopened_ref, reopened_cand = _save_reopen_both(ref, cand, fmt, tmp_path)
    assert [x.text for x in reopened_cand.paragraphs] == [x.text for x in reopened_ref.paragraphs]


@pytest.mark.common_api
@pytest.mark.parity
@pytest.mark.parametrize("style", [(True,False), (False,True), (True,True), (False,False)])
@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
def test_public_run_bold_italic_matches_python_docx(fmt, style, tmp_path):
    ref, cand = ReferenceDocument(), Document.new(fmt)
    for obj in (ref, cand):
        p = obj.add_paragraph()
        run = p.add_run("비교할 문자열")
        run.bold, run.italic = style
    rr, cc = _save_reopen_both(ref, cand, fmt, tmp_path)
    def styles(doc):
        return [(r.text, bool(r.bold), bool(r.italic)) for p in doc.paragraphs for r in p.runs if r.text]
    assert styles(cc) == styles(rr)


@pytest.mark.common_api
@pytest.mark.parity
@pytest.mark.parametrize("size", [(1,1), (2,3), (3,2)])
@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
def test_public_table_grid_matches_python_docx(fmt, size, tmp_path):
    ref, cand = ReferenceDocument(), Document.new(fmt)
    for obj in (ref, cand):
        table = obj.add_table(*size)
        for i in range(size[0]):
            for j in range(size[1]):
                table.cell(i, j).text = f"r{i}c{j}"
    rr, cc = _save_reopen_both(ref, cand, fmt, tmp_path)
    def values(doc):
        return [[c.text for c in row.cells] for row in doc.tables[0].rows]
    assert values(cc) == values(rr)


@pytest.mark.common_api
@pytest.mark.parity
@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
def test_public_table_add_row_matches_python_docx(fmt, tmp_path):
    ref, cand = ReferenceDocument(), Document.new(fmt)
    for obj in (ref, cand):
        tbl = obj.add_table(1, 2)
        tbl.cell(0, 0).text = "A"
        row = tbl.add_row()
        row.cells[1].text = "B"
    rr, cc = _save_reopen_both(ref, cand, fmt, tmp_path)
    assert [[c.text for c in r.cells] for r in cc.tables[0].rows] == [
        [c.text for c in r.cells] for r in rr.tables[0].rows
    ]


@pytest.mark.common_api
@pytest.mark.parity
@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
def test_public_metadata_title_and_author_matches_python_docx(fmt, tmp_path):
    ref, cand = ReferenceDocument(), Document.new(fmt)
    for obj in (ref, cand):
        obj.core_properties.title = "문서 제목"
        obj.core_properties.author = "Doc QA"
    rr, cc = _save_reopen_both(ref, cand, fmt, tmp_path)
    assert (cc.core_properties.title, cc.core_properties.author) == (
        rr.core_properties.title, rr.core_properties.author
    )
