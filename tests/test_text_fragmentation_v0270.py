from __future__ import annotations

import pytest

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from pywpdoc import Document, DocumentUnsupportedError
from pywpdoc.core.text import TextMap

XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _append_text(run_element, text: str, *, preserve: bool = False):
    node = OxmlElement("w:t")
    node.text = text
    if preserve:
        node.set(XML_SPACE, "preserve")
    run_element.append(node)
    return node


def _fragmented_single_run(paragraph, *parts: str, preserve: bool = False):
    paragraph.clear()
    run = paragraph.add_run("")
    run_element = run.native._r
    nodes = [_append_text(run_element, part, preserve=preserve) for part in parts]
    return run, nodes


def _insert_field_marker(paragraph) -> None:
    run = paragraph.add_run("").native._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run.append(begin)
    run.append(end)


def _insert_revision(paragraph, text: str = "IGNORED") -> None:
    revision = OxmlElement("w:ins")
    run = OxmlElement("w:r")
    _append_text(run, text)
    revision.append(run)
    paragraph.native._p.append(revision)


def _insert_sdt(paragraph, text: str = "IGNORED") -> None:
    sdt = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    _append_text(run, text)
    content.append(run)
    sdt.append(content)
    paragraph.native._p.append(sdt)


def test_docx_multiple_wt_nodes_are_separate_fragments_and_render_in_place() -> None:
    doc = Document.new("docx")
    paragraph = doc.add_paragraph("")
    run, nodes = _fragmented_single_run(paragraph, "{{cus", "tomer.name}}")
    run.bold = True

    text_map = TextMap.from_paragraph(paragraph)
    assert text_map.text == "{{customer.name}}"
    assert len(text_map.segments) == 2
    assert text_map.unsafe_boundaries == ()

    assert doc.render_template({"customer": {"name": "Alice"}}, everywhere=False) == 1
    assert paragraph.text == "Alice"
    assert nodes[0].text == "Alice"
    assert nodes[1].text in (None, "")
    assert run.bold is True
    assert len(run.native._r.findall(f"{{{W_NS}}}t")) == 2


def test_docx_soft_markers_and_hyperlink_do_not_break_one_template_tag() -> None:
    doc = Document.new("docx")
    paragraph = doc.add_paragraph("")
    paragraph.clear()
    paragraph.add_run("{{cus")

    proof = OxmlElement("w:proofErr")
    proof.set(qn("w:type"), "spellStart")
    paragraph.native._p.append(proof)

    hyperlink = OxmlElement("w:hyperlink")
    rel_id = paragraph.native.part.relate_to(
        "https://example.com", RT.HYPERLINK, is_external=True
    )
    hyperlink.set(qn("r:id"), rel_id)
    hyperlink_run = OxmlElement("w:r")
    _append_text(hyperlink_run, "tomer")
    hyperlink.append(hyperlink_run)
    paragraph.native._p.append(hyperlink)

    bookmark_start = OxmlElement("w:bookmarkStart")
    bookmark_start.set(qn("w:id"), "42")
    bookmark_start.set(qn("w:name"), "template")
    paragraph.native._p.append(bookmark_start)
    paragraph.add_run(".name}}")
    bookmark_end = OxmlElement("w:bookmarkEnd")
    bookmark_end.set(qn("w:id"), "42")
    paragraph.native._p.append(bookmark_end)

    text_map = TextMap.from_paragraph(paragraph)
    assert text_map.text == "{{customer.name}}"
    assert len(text_map.segments) == 3
    assert text_map.unsafe_boundaries == ()

    assert doc.render_template({"customer": {"name": "홍길동"}}, everywhere=False) == 1
    assert paragraph.text == "홍길동"
    assert paragraph.native._p.find(qn("w:hyperlink")) is not None
    assert paragraph.native._p.find(qn("w:bookmarkStart")) is not None
    assert paragraph.native._p.find(qn("w:proofErr")) is not None


def test_docx_xml_space_is_preserved_after_fragment_level_template_mutation() -> None:
    doc = Document.new("docx")
    paragraph = doc.add_paragraph("")
    _, nodes = _fragmented_single_run(paragraph, " {{na", "me}} ", preserve=True)

    assert TextMap.from_paragraph(paragraph).text == " {{name}} "
    assert doc.render_template({"name": "A"}, everywhere=False) == 1
    assert TextMap.from_paragraph(paragraph).text == " A "
    assert nodes[0].get(XML_SPACE) == "preserve"
    assert nodes[1].get(XML_SPACE) == "preserve"


@pytest.mark.parametrize("boundary", ["field", "revision", "sdt"])
def test_template_fails_closed_when_tag_crosses_opaque_docx_boundary(boundary: str) -> None:
    doc = Document.new("docx")
    paragraph = doc.add_paragraph("")
    paragraph.clear()
    paragraph.add_run("{{cus")
    if boundary == "field":
        _insert_field_marker(paragraph)
    elif boundary == "revision":
        _insert_revision(paragraph)
    else:
        _insert_sdt(paragraph)
    paragraph.add_run("tomer}}")

    text_map = TextMap.from_paragraph(paragraph)
    assert text_map.text == "{{customer}}"
    assert text_map.unsafe_boundaries

    with pytest.raises(DocumentUnsupportedError) as excinfo:
        doc.render_template({"customer": "X"}, everywhere=False)
    assert excinfo.value.code == "template-unsafe-text-boundary"


def test_find_and_replace_do_not_cross_opaque_boundary_and_limit_counts_safe_match() -> None:
    doc = Document.new("docx")
    unsafe = doc.add_paragraph("")
    unsafe.clear()
    unsafe.add_run("nee")
    _insert_field_marker(unsafe)
    unsafe.add_run("dle")
    safe = doc.add_paragraph("needle")

    matches = doc.find("needle", everywhere=False, count=1)
    assert len(matches) == 1
    assert matches[0].paragraph.text == safe.text

    assert doc.replace_text("needle", "X", everywhere=False, count=1) == 1
    assert TextMap.from_paragraph(unsafe).text == "needle"
    assert safe.text == "X"


def test_fragmented_tags_render_in_docx_table_header_and_footer() -> None:
    doc = Document.new("docx")

    table = doc.add_table(1, 1)
    cell_p = table.cell(0, 0).paragraphs[0]
    _fragmented_single_run(cell_p, "{{na", "me}}")

    header_p = doc.sections[0].header.paragraphs[0]
    _fragmented_single_run(header_p, "{{na", "me}}")

    footer_p = doc.sections[0].footer.paragraphs[0]
    _fragmented_single_run(footer_p, "{{na", "me}}")

    assert doc.render_template({"name": "Alice"}) == 3
    assert cell_p.text == "Alice"
    assert header_p.text == "Alice"
    assert footer_p.text == "Alice"


def test_tag_split_by_tab_or_break_is_not_treated_as_contiguous_text() -> None:
    doc = Document.new("docx")
    paragraph = doc.add_paragraph("")
    paragraph.clear()
    run = paragraph.add_run("")
    _append_text(run.native._r, "{{cus")
    run.native._r.append(OxmlElement("w:tab"))
    _append_text(run.native._r, "tomer}}")

    text_map = TextMap.from_paragraph(paragraph)
    assert text_map.text == "{{customer}}"
    assert text_map.unsafe_boundaries
    assert doc.find("{{customer}}", everywhere=False) == []
    assert doc.replace_text("{{customer}}", "X", everywhere=False) == 0
    with pytest.raises(DocumentUnsupportedError):
        doc.render_template({"customer": "X"}, everywhere=False)


def test_docx_fragmentation_survives_save_reopen_and_render(tmp_path) -> None:
    doc = Document.new("docx")
    paragraph = doc.add_paragraph("")
    paragraph.clear()
    paragraph.add_run("{{cus")

    proof = OxmlElement("w:proofErr")
    proof.set(qn("w:type"), "spellStart")
    paragraph.native._p.append(proof)

    hyperlink = OxmlElement("w:hyperlink")
    rel_id = paragraph.native.part.relate_to(
        "https://example.com", RT.HYPERLINK, is_external=True
    )
    hyperlink.set(qn("r:id"), rel_id)
    hyperlink_run = OxmlElement("w:r")
    _append_text(hyperlink_run, "tomer")
    hyperlink.append(hyperlink_run)
    paragraph.native._p.append(hyperlink)
    paragraph.add_run(".name}}")

    source = tmp_path / "fragmented.docx"
    doc.save(source)

    reopened = Document.open(source)
    assert reopened.render_template(
        {"customer": {"name": "Alice"}}, everywhere=False
    ) == 1
    assert reopened.paragraphs[-1].text == "Alice"

    out = tmp_path / "rendered.docx"
    reopened.save(out)
    final = Document.open(out)
    assert final.paragraphs[-1].text == "Alice"
