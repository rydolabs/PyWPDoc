"""Read-only DOCX/HWPX template validation, with legacy render regressions."""
from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

import pytest

from pywpdoc import Document, DocumentValueError
from pywpdoc.features.template import TemplateCheckIssue, TemplateCheckReport, check_template


def _parts(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as zf:
        return {name: hashlib.sha256(zf.read(name)).hexdigest() for name in zf.namelist()}


@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
def test_plain_variables_split_runs_and_return_dict(fmt):
    doc = Document.new(fmt)
    p = doc.add_paragraph("")
    p.clear()
    p.add_run("{{cus")
    p.add_run("tomer.name}} / {{missing}} / {{none}}")
    report = doc.check_template({"customer": {"name": "Kim"}, "none": None}, everywhere=False)
    assert isinstance(report, TemplateCheckReport)
    assert report.tag_count == 3
    assert report.checked_values == 3
    assert report.resolved_values == 2
    assert report.missing_keys == ("missing",)
    assert report.missing_count == 1
    assert report.issues[0].code == "template-value-missing"
    assert report.issues[0].tag == "{{missing}}"
    assert report.to_dict() == json.loads(json.dumps(report.to_dict()))
    assert isinstance(report.issues[0], TemplateCheckIssue)
    assert p.text == "{{customer.name}} / {{missing}} / {{none}}"


@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
def test_validation_and_rendering_remain_separate(fmt):
    doc = Document.new(fmt)
    p = doc.add_paragraph("Invoice {{id}}")
    assert doc.check_template({"id": 5}, everywhere=False).ok
    assert p.text == "Invoice {{id}}"
    assert doc.render_template({"id": 5}, everywhere=False) == 1
    assert p.text == "Invoice 5"
    assert doc.check_template({"id": 5}, everywhere=False).tag_count == 0


@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
def test_sections_only_validate_active_variables_and_parent_fallback(fmt):
    doc = Document.new(fmt)
    doc.add_paragraph("{{#items}}{{prefix}}:{{name}};{{/items}}{{^items}}{{unused}}{{/items}}")
    valid = doc.check_template({"prefix": "A", "items": [{"name": "x"}, {"name": "y"}]}, sections=True)
    assert valid.ok
    assert valid.checked_values == 4
    missing = doc.check_template({"prefix": "A", "items": [{"name": "x"}, {}]}, sections=True)
    assert missing.missing_keys == ("name",)
    assert missing.missing_count == 1
    hidden = doc.check_template({"items": []}, sections=True)
    assert hidden.missing_keys == ("unused",)


@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
def test_table_row_repeat_contexts_match_renderer(fmt):
    doc = Document.new(fmt)
    table = doc.add_table(2, 3)
    table.cell(1, 0).text = "{{#items}}{{name}}"
    table.cell(1, 1).text = "{{qty}}"
    table.cell(1, 2).text = "{{company}}{{/items}}"
    data = {"company": "ACME", "items": [{"name": "A", "qty": 1}, {"name": "B", "qty": 2}]}
    report = doc.check_template(data, sections=True, table_rows=True)
    assert report.ok, report.to_dict()
    assert report.tag_count == 5
    assert report.checked_values == 6
    assert [c.text for c in table.rows[1].cells] == ["{{#items}}{{name}}", "{{qty}}", "{{company}}{{/items}}"]
    bad = doc.check_template({"company": "ACME", "items": [{"name": "A"}]}, sections=True, table_rows=True)
    assert bad.missing_keys == ("qty",)
    assert doc.check_template({"items": []}, sections=True, table_rows=True).ok
    with pytest.raises(DocumentValueError) as exc:
        doc.check_template(data, sections=False, table_rows=True)
    assert exc.value.code == "template-row-repeat-requires-sections"


@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
def test_traversal_includes_table_header_footer_and_everywhere_filter(fmt):
    doc = Document.new(fmt)
    doc.add_paragraph("{{body}}")
    doc.add_table(1, 1).cell(0, 0).text = "{{cell}}"
    doc.sections[0].header.text = "{{header}}"
    doc.sections[0].footer.text = "{{footer}}"
    all_report = doc.check_template({"body": 1})
    assert set(all_report.missing_keys) == {"cell", "header", "footer"}
    short = doc.check_template({"body": 1}, everywhere=False)
    assert short.ok
    assert short.checked_values == 1


@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
@pytest.mark.parametrize("text,code", [
    ("{{name", "template-tag-unclosed"),
    ("{{name}}}}", "template-tag-orphan-close"),
    ("{{}}", "template-tag-empty"),
    ("{{{name}}}", "template-tag-invalid"),
    ("{{#items}}X{{/other}}", "template-section-mismatch"),
    ("{{#items}}X", "template-section-unbalanced"),
    ("{{>partial}}", "template-control-not-supported-v1"),
])
def test_malformed_tags_are_issues_not_mutations(fmt, text, code):
    doc = Document.new(fmt)
    p = doc.add_paragraph(text)
    report = doc.check_template({"name": "X", "items": True}, everywhere=False, sections=True)
    assert not report.ok
    assert code in [issue.code for issue in report.issues]
    assert p.text == text


@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
def test_section_requires_opt_in(fmt):
    doc = Document.new(fmt)
    doc.add_paragraph("{{#items}}{{.}}{{/items}}")
    report = doc.check_template({"items": [1]}, sections=False)
    assert not report.ok
    assert "template-control-not-supported-v1" in [i.code for i in report.issues]


def test_opaque_field_boundary_is_rejected_without_touching_the_docx():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document.new("docx")
    p = doc.add_paragraph("")
    p.clear()
    p.add_run("{{fi")
    run = p.add_run("").native._r
    for kind in ("begin", "end"):
        item = OxmlElement("w:fldChar")
        item.set(qn("w:fldCharType"), kind)
        run.append(item)
    p.add_run("eld}}")
    report = doc.check_template({"field": "x"}, everywhere=False)
    assert "template-unsafe-text-boundary" in [i.code for i in report.issues]
    assert p.text == "{{field}}"


@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
def test_package_parts_unchanged_by_preflight(fmt, tmp_path):
    doc = Document.new(fmt)
    doc.add_paragraph("{{#items}}{{name}}{{/items}} {{unknown}}")
    doc.add_table(1, 1).cell(0, 0).text = "{{other}}"
    before = tmp_path / f"before.{fmt}"
    after = tmp_path / f"after.{fmt}"
    doc.save(before)
    report = doc.check_template({"items": [{"name": "A"}]}, sections=True)
    assert not report.ok
    doc.save(after)
    assert _parts(before) == _parts(after)


@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
def test_malformed_in_inactive_branch_is_still_reported(fmt):
    doc = Document.new(fmt)
    doc.add_paragraph("{{#flag}}{{missing}} {{bad{{var}}{{/flag}}")
    report = doc.check_template({"flag": False}, sections=True)
    assert not report.ok
    assert "template-tag-invalid" in [i.code for i in report.issues]


@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
def test_nested_structural_rows_do_not_produce_false_pass(fmt):
    doc = Document.new(fmt)
    outer = doc.add_table(1, 2)
    outer.cell(0, 0).text = "{{#items}}{{name}}"
    inner = outer.cell(0, 1).add_table(1, 1)
    inner.cell(0, 0).text = "{{#children}}{{value}}{{/children}}"
    outer.cell(0, 1).paragraphs[-1].text = "{{/items}}"
    report = doc.check_template({"items": [{"name": "A", "children": [{"value": "ok"}]}]},
                                sections=True, table_rows=True)
    assert "template-nested-row-context-unverified" in [issue.code for issue in report.issues]


def test_backend_direct_check_and_pure_package_exports():
    doc = Document.new("hwpx")
    doc.add_paragraph("{{name}}")
    assert check_template(doc.backend, {"name": "A"}).ok
