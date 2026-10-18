from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

import pytest

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from pywpdoc import Document, DocumentUnsupportedError
from pywpdoc.features.template import render_section_text

ROOT = Path(__file__).resolve().parents[1]
MANIFEST = json.loads((ROOT / "tests/baseline/v0215_manifest.json").read_text(encoding="utf-8"))


def _zip_hashes(path: Path, selected: set[str] | None = None) -> dict[str, str]:
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        if selected is not None:
            names &= selected
        return {name: hashlib.sha256(zf.read(name)).hexdigest() for name in sorted(names)}


def _insert_field_marker(paragraph) -> None:
    run = paragraph.add_run("").native._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run.append(begin)
    run.append(end)


def test_pure_section_renderer_supports_list_context_parent_fallback_and_current_value() -> None:
    result = render_section_text(
        "{{#items}}{{prefix}}:{{name}}={{.}};{{/items}}",
        {
            "prefix": "P",
            "items": ["A", "B"],
        },
    )
    # Scalar items expose {{.}} while unresolved names can still fall back to parent.
    assert result.text == "P:{{name}}=A;P:{{name}}=B;"

    result = render_section_text(
        "{{#items}}{{prefix}}:{{name}};{{/items}}",
        {"prefix": "P", "items": [{"name": "A"}, {"name": "B"}]},
    )
    assert result.text == "P:A;P:B;"


def test_pure_section_renderer_supports_nested_and_inverted_sections() -> None:
    text = "{{#customer}}{{name}}:{{#items}}{{.}},{{/items}}{{^items}}empty{{/items}}{{/customer}}"
    assert render_section_text(
        text,
        {"customer": {"name": "A", "items": [1, 2]}},
    ).text == "A:1,2,"
    assert render_section_text(
        text,
        {"customer": {"name": "A", "items": []}},
    ).text == "A:empty"


@pytest.mark.parametrize("kind", ["hwpx", "docx"])
def test_inline_sections_are_opt_in_and_same_api_for_both_backends(kind: str) -> None:
    doc = Document.new(kind)
    paragraph = doc.add_paragraph("{{#items}}[{{name}}]{{/items}}{{^items}}EMPTY{{/items}}")

    changed = doc.render_template(
        {"items": [{"name": "A"}, {"name": "B"}]},
        everywhere=False,
        sections=True,
    )

    assert changed >= 2
    assert paragraph.text == "[A][B]"


@pytest.mark.parametrize("kind", ["hwpx", "docx"])
def test_inline_section_false_and_inverted_semantics(kind: str) -> None:
    doc = Document.new(kind)
    first = doc.add_paragraph("A{{#enabled}}YES{{/enabled}}B")
    second = doc.add_paragraph("A{{^items}}EMPTY{{/items}}B")

    doc.render_template(
        {"enabled": False, "items": []},
        everywhere=False,
        sections=True,
    )

    assert first.text == "AB"
    assert second.text == "AEMPTYB"


@pytest.mark.parametrize("kind", ["hwpx", "docx"])
def test_inline_sections_support_split_control_tags_and_variables_across_runs(kind: str) -> None:
    doc = Document.new(kind)
    paragraph = doc.add_paragraph("")
    paragraph.clear()
    first = paragraph.add_run("{{#it")
    first.bold = True
    paragraph.add_run("ems}}<{{na")
    last = paragraph.add_run("me}}>{{/items}}!")
    last.italic = True

    doc.render_template(
        {"items": [{"name": "A"}, {"name": "B"}]},
        everywhere=False,
        sections=True,
    )

    assert paragraph.text == "<A><B>!"
    # Inline structural expansion is intentionally a text-range mutation:
    # repeated content inherits the opening fragment's formatting while
    # surviving suffix content keeps its original run formatting.
    assert first.bold is True
    assert last.text == "!"
    assert last.italic is True


@pytest.mark.parametrize("kind", ["hwpx", "docx"])
def test_section_variables_can_fall_back_to_parent_context(kind: str) -> None:
    doc = Document.new(kind)
    paragraph = doc.add_paragraph("{{#customer}}{{company}}/{{name}}{{/customer}}")
    doc.render_template(
        {"company": "OpenAI", "customer": {"name": "Alice"}},
        everywhere=False,
        sections=True,
    )
    assert paragraph.text == "OpenAI/Alice"


def test_sections_remain_fail_closed_by_default_for_backward_compatibility() -> None:
    doc = Document.new("hwpx")
    doc.add_paragraph("{{#items}}{{.}}{{/items}}")
    with pytest.raises(DocumentUnsupportedError) as excinfo:
        doc.render_template({"items": [1, 2]}, everywhere=False)
    assert excinfo.value.code == "template-control-not-supported-v1"


def test_docx_section_crossing_opaque_boundary_fails_safe() -> None:
    doc = Document.new("docx")
    paragraph = doc.add_paragraph("")
    paragraph.clear()
    paragraph.add_run("{{#enabled}}A")
    _insert_field_marker(paragraph)
    paragraph.add_run("B{{/enabled}}")

    with pytest.raises(DocumentUnsupportedError) as excinfo:
        doc.render_template(
            {"enabled": True},
            everywhere=False,
            sections=True,
        )
    assert excinfo.value.code == "template-unsafe-text-boundary"


@pytest.mark.parametrize("kind", ["hwpx", "docx"])
def test_sections_render_inside_table_header_and_footer(kind: str) -> None:
    doc = Document.new(kind)
    table = doc.add_table(1, 1)
    table.cell(0, 0).text = "{{#show}}CELL{{/show}}"
    doc.sections[0].header.text = "{{#show}}HEAD{{/show}}"
    doc.sections[0].footer.text = "{{^show}}NO{{/show}}"

    doc.render_template({"show": True}, sections=True)

    assert table.cell(0, 0).text == "CELL"
    assert "HEAD" in doc.sections[0].header.text
    assert "NO" not in doc.sections[0].footer.text


def test_section_feature_keeps_sibling_feature_dependency_rule() -> None:
    source = (ROOT / "src/pywpdoc/features/template/sections.py").read_text(encoding="utf-8")
    assert "features.find" not in source
    assert "features.replace" not in source
    assert "pywpdoc.hwpx" not in source
    assert "pywpdoc.docx" not in source
    assert "python_docx" not in source


def test_sections_true_noop_preserves_v0215_hwpx_golden(tmp_path: Path) -> None:
    source = ROOT / MANIFEST["source_fixture"]
    doc = Document.open(source)
    assert doc.render_template({"unused": "X"}, sections=True) == 0
    out = tmp_path / "noop-section.hwpx"
    doc.save(out)
    assert _zip_hashes(out) == MANIFEST["hwpx_package_parts"]


def test_sections_true_noop_preserves_docx_selected_parts(tmp_path: Path) -> None:
    source = ROOT / MANIFEST["source_fixture"]
    selected = set(MANIFEST["docx_selected_parts"])

    baseline = Document.rebuild(source, format="docx", preserve_layout=True)
    before = tmp_path / "before.docx"
    baseline.save(before)

    changed = Document.rebuild(source, format="docx", preserve_layout=True)
    assert changed.render_template({"unused": "X"}, sections=True) == 0
    after = tmp_path / "after.docx"
    changed.save(after)

    assert _zip_hashes(after, selected) == _zip_hashes(before, selected)
    assert _zip_hashes(after, selected) == MANIFEST["docx_selected_parts"]


def test_section_generated_braces_are_data_not_recursive_template_syntax() -> None:
    doc = Document.new("docx")
    paragraph = doc.add_paragraph("{{#items}}{{value}}{{/items}} / {{secret}}")
    doc.render_template(
        {"items": [{"value": "{{secret}}"}], "secret": "ROOT"},
        everywhere=False,
        sections=True,
    )
    assert paragraph.text == "{{secret}} / ROOT"


def test_nested_section_generated_braces_are_not_reparsed() -> None:
    doc = Document.new("docx")
    paragraph = doc.add_paragraph(
        "{{#outer}}{{#items}}{{value}}{{/items}}{{/outer}} / {{secret}}"
    )
    doc.render_template(
        {
            "outer": {"items": [{"value": "{{secret}}"}]},
            "secret": "ROOT",
        },
        everywhere=False,
        sections=True,
    )
    assert paragraph.text == "{{secret}} / ROOT"
