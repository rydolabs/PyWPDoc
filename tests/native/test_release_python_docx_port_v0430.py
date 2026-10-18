"""Release candidate: semantic port of representative python-docx scenarios.

The upstream python-docx tests contain package/XML internals and mocks that do
not belong to the PyWPDoc Common API. These tests port public behavior from
python-docx quickstart + document/paragraph/table scenarios, not verbatim code.

Reference: https://python-docx.readthedocs.io/en/latest/user/quickstart.html
          https://github.com/python-openxml/python-docx/blob/master/tests/test_document.py
          https://github.com/python-openxml/python-docx/blob/master/tests/test_table.py
"""
from __future__ import annotations
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document as NativeDocx
from docx.enum.section import WD_ORIENT as NativeOrient
from docx.shared import Mm as NativeMm
from pywpdoc import Document, Mm, WD_ORIENT


def _saved_common(fmt, doc, tmp_path, stem):
    path = tmp_path / (stem + '.' + fmt)
    doc.save(path)
    with ZipFile(path) as archive:
        assert archive.testzip() is None
    return Document.open(path)


def _saved_reference(doc, tmp_path, stem):
    path = tmp_path / (stem + '.docx')
    doc.save(path)
    return NativeDocx(path)


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_port_add_paragraph_and_runs(fmt, tmp_path):
    baseline = NativeDocx()
    baseline.add_paragraph('기본 문단')
    bp = baseline.add_paragraph()
    bp.add_run('A\tB\nC')
    candidate = Document.new(fmt)
    candidate.add_paragraph('기본 문단')
    cp = candidate.add_paragraph()
    cp.add_run('A\tB\nC')
    ref = _saved_reference(baseline, tmp_path, 'ref_para')
    got = _saved_common(fmt, candidate, tmp_path, 'common_para')
    assert [p.text for p in got.paragraphs] == [p.text for p in ref.paragraphs]


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_port_formatted_runs(fmt, tmp_path):
    baseline = NativeDocx()
    p = baseline.add_paragraph()
    p.add_run('BOLD').bold = True
    p.add_run(' italic').italic = True
    candidate = Document.new(fmt)
    p = candidate.add_paragraph()
    p.add_run('BOLD').bold = True
    p.add_run(' italic').italic = True
    ref = _saved_reference(baseline, tmp_path, 'ref_run')
    got = _saved_common(fmt, candidate, tmp_path, 'common_run')
    # HWPX serializes an extra empty run and resolves unspecified flags to False.
    # Compare visible text and the effective formatting, not tri-state internals.
    common_runs = [(r.text, bool(r.bold), bool(r.italic)) for r in got.paragraphs[-1].runs if r.text]
    native_runs = [(r.text, bool(r.bold), bool(r.italic)) for r in ref.paragraphs[-1].runs if r.text]
    assert common_runs == native_runs


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_port_insert_paragraph_before(fmt, tmp_path):
    baseline, candidate = NativeDocx(), Document.new(fmt)
    baseline.add_paragraph('second').insert_paragraph_before('first')
    candidate.add_paragraph('second').insert_paragraph_before('first')
    ref = _saved_reference(baseline, tmp_path, 'ref_insert')
    got = _saved_common(fmt, candidate, tmp_path, 'common_insert')
    assert [p.text for p in got.paragraphs] == [p.text for p in ref.paragraphs]


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_port_table_text_grid(fmt, tmp_path):
    baseline, candidate = NativeDocx(), Document.new(fmt)
    for doc in (baseline, candidate):
        table = doc.add_table(rows=2, cols=2)
        for row in range(2):
            for col in range(2):
                table.cell(row, col).text = f'R{row}C{col}'
    ref = _saved_reference(baseline, tmp_path, 'ref_grid')
    got = _saved_common(fmt, candidate, tmp_path, 'common_grid')
    assert [[c.text for c in row.cells] for row in got.tables[0].rows] == [
        [c.text for c in row.cells] for row in ref.tables[0].rows
    ]


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_port_merged_cells_logical_grid(fmt, tmp_path):
    baseline, candidate = NativeDocx(), Document.new(fmt)
    for doc in (baseline, candidate):
        table = doc.add_table(2, 2)
        table.cell(0, 0).text = 'merged'
        table.cell(0, 0).merge(table.cell(1, 1))
    ref = _saved_reference(baseline, tmp_path, 'ref_merge')
    got = _saved_common(fmt, candidate, tmp_path, 'common_merge')
    assert [[c.text for c in row.cells] for row in got.tables[0].rows] == [
        [c.text for c in row.cells] for row in ref.tables[0].rows
    ]


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_port_nested_table_visibility(fmt, tmp_path):
    baseline, candidate = NativeDocx(), Document.new(fmt)
    for doc in (baseline, candidate):
        outer = doc.add_table(1, 1)
        outer.cell(0, 0).text = 'outer'
        outer.cell(0, 0).add_table(1, 1).cell(0, 0).text = 'inner'
    ref = _saved_reference(baseline, tmp_path, 'ref_nested')
    got = _saved_common(fmt, candidate, tmp_path, 'common_nested')
    assert len(ref.tables) == len(got.tables) == 1
    assert len(got.tables[0].cell(0, 0).tables) == len(ref.tables[0].cell(0, 0).tables) == 1
    assert got.tables[0].cell(0, 0).tables[0].cell(0, 0).text == 'inner'


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_port_core_properties(fmt, tmp_path):
    baseline, candidate = NativeDocx(), Document.new(fmt)
    for doc in (baseline, candidate):
        doc.core_properties.title = 'Ported Contract'
        doc.core_properties.author = 'QA'
        doc.core_properties.subject = 'Compatibility'
    ref = _saved_reference(baseline, tmp_path, 'ref_props')
    got = _saved_common(fmt, candidate, tmp_path, 'common_props')
    for field in ('title', 'author', 'subject'):
        assert getattr(got.core_properties, field) == getattr(ref.core_properties, field)


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_port_section_margin_orientation(fmt, tmp_path):
    baseline, candidate = NativeDocx(), Document.new(fmt)
    baseline.sections[0].left_margin = NativeMm(18)
    baseline.sections[0].orientation = NativeOrient.LANDSCAPE
    candidate.sections[0].left_margin = Mm(18)
    candidate.sections[0].orientation = WD_ORIENT.LANDSCAPE
    ref = _saved_reference(baseline, tmp_path, 'ref_section')
    got = _saved_common(fmt, candidate, tmp_path, 'common_section')
    assert abs(got.sections[0].left_margin.mm - ref.sections[0].left_margin.mm) < 0.06
    assert got.sections[0].orientation == WD_ORIENT.LANDSCAPE


def test_port_requires_format_specific_escape_for_upstream_comments():
    # Intentionally NOT a common API test: python-docx's add_comment() is not
    # a declared common contract. No fabricated parity or silent approximation.
    candidate = Document.new('docx')
    # The thin facade delegates DOCX-only methods to its backend, but the
    # same operation is unavailable in HWPX: not portable common behavior.
    assert hasattr(candidate, 'add_comment')
    assert not hasattr(Document.new('hwpx'), 'add_comment')

@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_port_table_add_row(fmt, tmp_path):
    baseline, candidate = NativeDocx(), Document.new(fmt)
    for doc in (baseline, candidate):
        table = doc.add_table(1, 2)
        table.cell(0, 0).text = 'header'
        new_row = table.add_row()
        new_row.cells[0].text = 'new'
        new_row.cells[1].text = 'row'
    ref = _saved_reference(baseline, tmp_path, 'ref_addrow')
    got = _saved_common(fmt, candidate, tmp_path, 'common_addrow')
    assert [[c.text for c in r.cells] for r in got.tables[0].rows] == [
        [c.text for c in r.cells] for r in ref.tables[0].rows
    ]


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_port_paragraph_clear(fmt, tmp_path):
    baseline, candidate = NativeDocx(), Document.new(fmt)
    for doc in (baseline, candidate):
        p = doc.add_paragraph()
        p.add_run('discard')
        p.add_run(' me')
        p.clear()
        p.add_run('kept')
    ref = _saved_reference(baseline, tmp_path, 'ref_clear')
    got = _saved_common(fmt, candidate, tmp_path, 'common_clear')
    assert got.paragraphs[-1].text == ref.paragraphs[-1].text == 'kept'


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_port_next_paragraph_style(fmt, tmp_path):
    from docx.enum.style import WD_STYLE_TYPE as NativeStyle
    from pywpdoc import WD_STYLE_TYPE
    baseline, candidate = NativeDocx(), Document.new(fmt)
    a = baseline.styles.add_style('Port Question', NativeStyle.PARAGRAPH)
    b = baseline.styles.add_style('Port Answer', NativeStyle.PARAGRAPH)
    a.next_paragraph_style = b
    if fmt == 'hwpx':
        a = candidate.styles.add_style('Port Question', WD_STYLE_TYPE.PARAGRAPH, clone_from='Body')
        b = candidate.styles.add_style('Port Answer', WD_STYLE_TYPE.PARAGRAPH, clone_from='Body')
    else:
        a = candidate.styles.add_style('Port Question', WD_STYLE_TYPE.PARAGRAPH)
        b = candidate.styles.add_style('Port Answer', WD_STYLE_TYPE.PARAGRAPH)
    a.next_paragraph_style = b
    ref = _saved_reference(baseline, tmp_path, 'ref_style')
    got = _saved_common(fmt, candidate, tmp_path, 'common_style')
    assert got.styles['Port Question'].next_paragraph_style.name == ref.styles['Port Question'].next_paragraph_style.name
