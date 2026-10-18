# -*- coding: utf-8 -*-
"""Long/stress document differential QA: python-docx vs PyWPDoc Common API.

Checks semantics after save/reopen. Real page count and visual layout require a
renderer/Word/Hangul and are intentionally never inferred from paragraph count.
No product source files are changed.
"""
from __future__ import annotations
import argparse
from dataclasses import dataclass
import hashlib
import io
import json
from pathlib import Path
import shutil
import struct
import time
from zipfile import ZipFile
import zlib

from docx import Document as NativeDocument
from docx.shared import Mm as NativeMm
from pywpdoc import Document, Mm


PROFILES = {
    'long': {'paragraphs': 520, 'tables': 8, 'rows_per_table': 10, 'images': 2},
    'stress': {'paragraphs': 1100, 'tables': 12, 'rows_per_table': 18, 'images': 3},
}


def png_fixture(index: int) -> bytes:
    """Distinct 240x70 PNGs from the standard library; no Pillow needed."""
    width,height=240,70
    rgb=bytearray()
    for y in range(height):
        rgb.append(0)  # filter byte
        for x in range(width):
            stripe=(x // 32 + index) % 4
            color=((25 + stripe*55)%256, (70+index*27+y)%256, (x+y*2)%256)
            rgb.extend(color)
    def chunk(tag,content):
        return struct.pack('>I',len(content))+tag+content+struct.pack('>I',zlib.crc32(tag+content)&0xffffffff)
    return b'\x89PNG\r\n\x1a\n'+chunk(b'IHDR',struct.pack('>IIBBBBB',width,height,8,2,0,0,0))+chunk(b'IDAT',zlib.compress(bytes(rgb),6))+chunk(b'IEND',b'')


def payload_line(idx: int) -> str:
    # Realistic Korean and Latin mixed text; deterministic IDs in every paragraph.
    fields=[
        '문서 처리의 정확성을 확인하기 위해 본문, 단락 경계, 표 식별자와 숫자 데이터를 함께 비교합니다.',
        '한국어 English 2026-10-14 / QA-{:05d} / 전표번호 A{:04d}를 추적합니다.'.format(idx,idx),
        '서식이 다른 여러 run 사이에서도 순서가 유지되고 저장·재열기 이후 텍스트가 유실되지 않아야 합니다.',
    ]
    return f'[PARA-{idx:05d}] ' + ' '.join(fields)


def add_paragraph(doc, text, idx):
    p=doc.add_paragraph()
    # Same public operations for python-docx and PyWPDoc.
    third=max(1,len(text)//3)
    a=p.add_run(text[:third]); a.bold=(idx%7==0)
    b=p.add_run(text[third:third*2]); b.italic=(idx%11==0)
    p.add_run(text[third*2:])
    return p


def build(profile: str, backend: str, out: Path):
    cfg=PROFILES[profile]
    is_reference=backend=='reference'
    doc=NativeDocument() if is_reference else Document.new(backend)
    # The factory templates have different DEFAULT paper sizes (US Letter
    # for python-docx versus A4 for PyWPDoc). Explicitly equalize geometry
    # before testing page-level rendering and never claim default parity.
    section=doc.sections[0]
    mm=NativeMm if is_reference else Mm
    for prop,value in {'page_width':210,'page_height':297,'left_margin':18,
                       'right_margin':20,'top_margin':20,'bottom_margin':20,
                       'header_distance':15,'footer_distance':15}.items():
        setattr(section,prop,mm(value))
    doc.core_properties.title=f'PyWPDoc differential {profile} profile'
    doc.core_properties.author='PyWPDoc QA'
    doc.core_properties.subject='Long-document compatibility stress probe'
    heading=f'PyWPDoc LONG DOCUMENT QA — {profile.upper()}'
    add_paragraph(doc,heading,0)
    add_paragraph(doc,'비교 기준: 문단 순서, run의 실제 표시 서식, 표·이미지·텍스트, 저장·재열기',1)
    table_at=set(round((x+1)*cfg['paragraphs']/(cfg['tables']+1)) for x in range(cfg['tables']))
    image_at=set(round((x+1)*cfg['paragraphs']/(cfg['images']+1)) for x in range(cfg['images']))
    table_index=0; image_index=0
    for i in range(cfg['paragraphs']):
        add_paragraph(doc, payload_line(i),i+2)
        if i in table_at:
            table_index+=1
            tbl=doc.add_table(cfg['rows_per_table'],3)
            for r in range(cfg['rows_per_table']):
                for c in range(3):tbl.cell(r,c).text=f'T{table_index:02d}-R{r:03d}-C{c} QA 데이터 {r*3+c:04d}'
        if i in image_at:
            image_index+=1
            raw=png_fixture(image_index)
            if is_reference: doc.add_picture(io.BytesIO(raw), width=NativeMm(34),height=NativeMm(10))
            else: doc.add_picture(raw,'png',width=Mm(34),height=Mm(10))
    # Nested table appears in python-docx feature/table tests. Exclude from
    # top-level table count but compare its semantic text separately.
    nested=doc.add_table(1,1)
    nested.cell(0,0).text='NESTED-OUTER'
    nested.cell(0,0).add_table(1,1).cell(0,0).text='NESTED-INNER'
    add_paragraph(doc, f'[FINAL-{profile.upper()}] CHECKSUM TAIL',cfg['paragraphs']+2)
    start=time.monotonic();doc.save(out)
    return {'seconds_save':round(time.monotonic()-start,3),'bytes':out.stat().st_size,
        'paragraphs_generated':cfg['paragraphs']+3, 'top_level_tables_generated':cfg['tables']+1,
        'images_generated':cfg['images'],'sha256':hashlib.sha256(out.read_bytes()).hexdigest()}


def package_check(path:Path):
    from lxml import etree
    with ZipFile(path) as z:
        assert z.testzip() is None, f'corrupt ZIP: {path.name}'
        names=z.namelist()
        for name in names:
            if name.lower().endswith(('.xml','.rels','.hpf')):
                etree.fromstring(z.read(name),etree.XMLParser(resolve_entities=False,recover=False))
        return {'package_members':len(names),'package_sha256':hashlib.sha256(path.read_bytes()).hexdigest()}


def structure(doc):
    pars=[p.text for p in doc.paragraphs if p.text]
    # Compare bold formatting on visible characters, not DOCX run boundaries.
    # MS Word is free to split a run into many text-equivalent runs on save.
    # Keep paragraph index and contiguous bold text, while ignoring empty runs.
    # Here the fixture sets bold explicitly; unspecified and False both mean
    # "not bold" for this test. Never use this simplification for inherited
    # style-based formatting without checking the effective style.
    formatted=[]
    for paragraph_index, paragraph in enumerate(doc.paragraphs):
        previous_bold=False
        chunk=''
        for run in paragraph.runs:
            if not run.text:
                continue
            is_bold=bool(run.bold)
            if is_bold:
                if previous_bold:
                    chunk += run.text
                else:
                    chunk = run.text
            elif previous_bold:
                formatted.append((paragraph_index, chunk))
                chunk=''
            previous_bold=is_bold
        if previous_bold:
            formatted.append((paragraph_index, chunk))
    # For a parent cell containing a nested table, DOCX and HWPX may append
    # a different number of required empty terminal paragraphs. Strip only
    # terminal newlines, retaining interior text and line breaks.
    grids=[[[cell.text.rstrip('\n') if cell.tables else cell.text for cell in r.cells] for r in t.rows] for t in doc.tables]
    nest=doc.tables[-1].cell(0,0).tables[0].cell(0,0).text
    return {'paragraphs':pars,'grids':grids,'nested_text':nest,
        'image_count':len(doc.pictures) if hasattr(doc,'pictures') else len(doc.inline_shapes),
        'title':doc.core_properties.title,'author':doc.core_properties.author,
        'bold_snippets':formatted}


def run(profile: str, root: Path):
    root.mkdir(parents=True,exist_ok=True)
    result={'profile':profile,'format_results':{},'checks':{},'notes':[]}
    errors=[]
    try:
        # Save all source cases independently; a failure in HWPX must not
        # prevent the reference DOCX from existing in output.
        files={}
        for backend,ext in [('reference','docx'),('docx','docx'),('hwpx','hwpx')]:
            path=root/f'{profile}_{backend}.{ext}'
            try:
                timing=build(profile,backend,path)
                inspect=package_check(path)
                result['format_results'][backend]={'status':'PASS','file':path.name,**timing,**inspect}
                files[backend]=path
            except Exception as exc:
                result['format_results'][backend]={'status':'ERROR','error':repr(exc)};errors.append(f'{backend}: {exc!r}')
        fingerprints={}
        for backend,path in files.items():
            try:
                started=time.monotonic();loaded=NativeDocument(path) if backend=='reference' else Document.open(path)
                fingerprints[backend]=structure(loaded)
                result['format_results'][backend]['seconds_reopen_scan']=round(time.monotonic()-started,3)
                result['format_results'][backend]['visible_paragraph_count']=len(fingerprints[backend]['paragraphs'])
            except Exception as exc:
                result['format_results'][backend]['status']='ERROR';result['format_results'][backend]['read_error']=repr(exc)
                errors.append(f'{backend} read: {exc!r}')
        for backend in ('docx','hwpx'):
            if backend not in fingerprints or 'reference' not in fingerprints:
                result['checks'][backend]={'status':'SKIP','reason':'reference/candidate unavailable'}
                continue
            target=fingerprints[backend];reference=fingerprints['reference']
            tests={k:target[k]==reference[k] for k in ('paragraphs','grids','nested_text','image_count','title','author','bold_snippets')}
            result['checks'][backend]={'status':'PASS' if all(tests.values()) else 'FAIL','details':tests,
                 'reference_paragraph_count':len(reference['paragraphs']),
                 'candidate_paragraph_count':len(target['paragraphs']),
                 'reference_table_count':len(reference['grids']),'candidate_table_count':len(target['grids'])}
        if errors: result['notes']=errors
        result['overall']='PASS' if all(x.get('status')=='PASS' for x in result['checks'].values()) and not errors else 'FAIL'
    except Exception as exc:
        result['overall']='ERROR';result['notes'].append(repr(exc))
    (root/f'{profile}_summary.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return result


def main():
    p=argparse.ArgumentParser()
    p.add_argument('--profile',choices=['long','stress','both'],default='long')
    p.add_argument('--output-dir',type=Path,default=Path('long_qa_results'))
    args=p.parse_args()
    profiles=list(PROFILES) if args.profile=='both' else [args.profile]
    output=[]
    for profile in profiles:
        print(f'[RUN] {profile}',flush=True)
        report=run(profile,args.output_dir)
        print(f'[DONE] {profile}: {report["overall"]}; checks={report["checks"]}',flush=True)
        output.append(report)
    (args.output_dir/'all_profiles.json').write_text(json.dumps(output,ensure_ascii=False,indent=2),encoding='utf-8')
    return 0 if all(x['overall']=='PASS' for x in output) else 1

if __name__=='__main__':raise SystemExit(main())
