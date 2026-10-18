"""Regression tests for Word's run splitting and actual formatting loss."""
from pathlib import Path
import pytest
from docx import Document as DocxDocument
from extended_document_qa import structure


def build_doc(path:Path, parts):
    d=DocxDocument()
    p=d.add_paragraph()
    for text,bold in parts:
        r=p.add_run(text)
        if bold is not None: r.bold=bold
    d.add_table(1,1).cell(0,0).text='TABLE'
    # structure() expects a final nested table
    d.tables[-1].cell(0,0).add_table(1,1).cell(0,0).text='NESTED-INNER'
    d.save(path)
    return path


def test_word_split_same_semantics(tmp_path):
    first=build_doc(tmp_path/'a.docx', [('[PARA-00005] 한국어 English 2026',True), (' END',None)])
    second=build_doc(tmp_path/'b.docx', [('[PARA-00005] ',True),('한국어',True),(' English',True),(' 2026',True),(' END',None)])
    assert structure(DocxDocument(first))['bold_snippets'] == structure(DocxDocument(second))['bold_snippets']


def test_actual_loss_is_detected(tmp_path):
    first=build_doc(tmp_path/'a.docx', [('[PARA-00005] 한국어 English',True),(' END',None)])
    second=build_doc(tmp_path/'b.docx', [('[PARA-00005] 한국어',True),(' English',False),(' END',None)])
    assert structure(DocxDocument(first))['bold_snippets'] != structure(DocxDocument(second))['bold_snippets']


def test_bold_moved_between_paragraphs_is_detected(tmp_path):
    def build(path,position):
        d=DocxDocument()
        for i in range(2):
            r=d.add_paragraph().add_run('same text')
            r.bold=(i==position)
        d.add_table(1,1).cell(0,0).add_table(1,1).cell(0,0).text='NESTED-INNER'
        d.save(path)
        return path
    a=build(tmp_path/'a.docx',0)
    b=build(tmp_path/'b.docx',1)
    assert structure(DocxDocument(a))['bold_snippets'] != structure(DocxDocument(b))['bold_snippets']


def test_empty_run_is_ignored(tmp_path):
    a=build_doc(tmp_path/'a.docx', [('ab',True)])
    b=build_doc(tmp_path/'b.docx', [('a',True),('',False),('b',True)])
    assert structure(DocxDocument(a))['bold_snippets'] == structure(DocxDocument(b))['bold_snippets']
