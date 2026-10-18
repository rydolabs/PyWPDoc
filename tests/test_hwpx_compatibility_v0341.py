"""v0.34.1 HWPX creation compatibility and font-size unit guards."""
from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

import pytest
from lxml import etree
from docx.shared import Pt as DocxPt

from pywpdoc import Document, DocumentValueError
from pywpdoc._engine import Document as NativeHwpxDocument, Pt, validate_hwpx_file
from pywpdoc._engine.errors import HwpxValueError


def _version(path: Path) -> dict[str, str]:
    with ZipFile(path) as zf:
        root = etree.fromstring(zf.read("version.xml"))
    return dict(root.attrib)


def _version_bytes(path: Path) -> bytes:
    with ZipFile(path) as zf:
        return zf.read("version.xml")


def _first_user_char_height(path: Path, text: str) -> int:
    hp = "http://www.hancom.co.kr/hwpml/2011/paragraph"
    hh = "http://www.hancom.co.kr/hwpml/2011/head"
    ns = {"hp": hp, "hh": hh}
    with ZipFile(path) as zf:
        section = etree.fromstring(zf.read("Contents/section0.xml"))
        header = etree.fromstring(zf.read("Contents/header.xml"))
    runs = section.xpath(f".//hp:p[.//hp:t[contains(., {text!r})]]/hp:run[hp:t]", namespaces=ns)
    assert runs
    ref = runs[0].get("charPrIDRef")
    props = header.xpath(f".//hh:charPr[@id={ref!r}]", namespaces=ns)
    assert len(props) == 1
    return int(props[0].get("height"))


def test_new_hwpx_defaults_to_real_hancom2020_profile(tmp_path: Path) -> None:
    doc = Document.new("hwpx")
    doc.add_paragraph("한글 2020 호환")
    out = tmp_path / "default-2020.hwpx"
    doc.save(out)

    attrs = _version(out)
    assert attrs["major"] == "5"
    assert attrs["minor"] == "1"
    assert attrs["micro"] == "0"
    assert attrs["buildNumber"] == "1"
    assert attrs["xmlVersion"] == "1.2"
    assert attrs["appVersion"].startswith("11, 0, 0, 2129")
    assert validate_hwpx_file(out).ok
    assert Document.open(out).paragraphs[-1].text == "한글 2020 호환"


def test_native_hwpx_latest_profile_preserves_previous_13x_skeleton(tmp_path: Path) -> None:
    doc = NativeHwpxDocument.new(compatibility="latest")
    doc.add_paragraph("latest")
    out = tmp_path / "latest.hwpx"
    doc.save(out)
    attrs = _version(out)
    assert attrs["xmlVersion"] == "1.5"
    assert attrs["appVersion"].startswith("13, 0, 0, 1408")


def test_invalid_hwpx_compatibility_profile_fails_closed() -> None:
    with pytest.raises(HwpxValueError) as excinfo:
        NativeHwpxDocument.new(compatibility="hancom2014")
    assert excinfo.value.code == "hwpx-compatibility-profile-invalid"

    with pytest.raises(DocumentValueError) as common_exc:
        Document.new("docx", compatibility="hancom2020")
    assert common_exc.value.code == "common-compatibility-not-applicable"


def test_open_edit_save_preserves_original_version_xml_byte_exact(tmp_path: Path) -> None:
    source = Path(__file__).parent / "fixtures" / "original_notice.hwpx"
    before = _version_bytes(source)
    doc = Document.open(source)
    doc.add_paragraph("version preservation")
    out = tmp_path / "edited.hwpx"
    doc.save(out)
    assert _version_bytes(out) == before


def test_compare_review_preserves_new_document_version_xml(tmp_path: Path) -> None:
    old = Document.new("hwpx", compatibility="hancom2020")
    new = Document.new("hwpx", compatibility="hancom2020")
    old.add_paragraph("before")
    new.add_paragraph("after")
    new_path = tmp_path / "new.hwpx"
    new.save(new_path)
    review = old.compare(new, author="Compat QA", date="2026-10-09T01:00:00Z", everywhere=False)
    review_path = tmp_path / "review.hwpx"
    review.save(review_path)
    assert _version_bytes(review_path) == _version_bytes(new_path)


def test_hwpx_font_size_accepts_python_docx_length_by_physical_points(tmp_path: Path) -> None:
    doc = Document.new("hwpx")
    run = doc.add_paragraph("foreign Pt").runs[0]
    run.font.size = DocxPt(14)
    assert run.font.size is not None
    assert run.font.size.pt == pytest.approx(14.0, abs=0.01)

    out = tmp_path / "foreign-pt.hwpx"
    doc.save(out)
    assert _first_user_char_height(out, "foreign Pt") == 1400
    reopened = Document.open(out)
    assert reopened.paragraphs[-1].runs[0].font.size.pt == pytest.approx(14.0, abs=0.01)


def test_hwpx_font_size_rejects_nonpositive_or_nonnumeric_values() -> None:
    doc = Document.new("hwpx")
    run = doc.add_paragraph("bad size").runs[0]
    for value in (0, -1, float("nan"), True, object()):
        with pytest.raises(HwpxValueError) as excinfo:
            run.font.size = value  # type: ignore[assignment]
        assert excinfo.value.code == "font-size-invalid"


def test_native_pywpdoc_pt_contract_still_uses_hwpunit_storage() -> None:
    doc = NativeHwpxDocument.new()
    run = doc.add_paragraph("native Pt").runs[0]
    run.font.size = Pt(13.5)
    assert run.font.size == Pt(13.5)
