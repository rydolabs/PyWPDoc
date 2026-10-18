"""v0.43.1rc2 safety regressions for untracked changes in body-only compare."""
from __future__ import annotations

from io import BytesIO
from hashlib import sha256

import pytest
from lxml import etree
from pywpdoc import Document
from pywpdoc._engine.core.errors import DocumentUnsupportedError


ERROR_CODE = "compare-outside-scope-changes-unsupported"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _docx_pair():
    a, b = Document.new("docx"), Document.new("docx")
    for d in (a, b):
        d.add_paragraph("identical body")
        d.add_table(1, 2).cell(0, 0).text = "unchanged table"
    return a, b


def _fail_closed(a, b):
    with pytest.raises(DocumentUnsupportedError) as exc:
        a.compare(b, everywhere=False)
    assert exc.value.code == ERROR_CODE


def _table_property(table, local):
    tbl = table.native._tbl
    pr = tbl.tblPr
    element = etree.Element("{" + W + "}" + local)
    pr.append(element)
    return element


def test_docx_table_only_border_change_fails_closed():
    a, b = _docx_pair()
    _table_property(b.tables[0], "tblBorders").append(
        etree.Element("{" + W + "}top", {"{" + W + "}val": "single"})
    )
    prior = etree.tostring(a.tables[0].native._tbl)
    _fail_closed(a, b)
    assert etree.tostring(a.tables[0].native._tbl) == prior
    assert b.tables[0].native._tbl.xpath(".//w:tblBorders")


def test_docx_cell_only_shading_change_fails_closed():
    a, b = _docx_pair()
    pr = b.tables[0].cell(0, 0).native._tc.get_or_add_tcPr()
    pr.append(etree.Element("{" + W + "}shd", {"{" + W + "}fill": "FFF000"}))
    _fail_closed(a, b)


def test_docx_cell_only_border_change_fails_closed():
    a, b = _docx_pair()
    pr = b.tables[0].cell(0, 0).native._tc.get_or_add_tcPr()
    pr.append(etree.Element("{" + W + "}tcBorders"))
    _fail_closed(a, b)


def test_docx_cell_merge_without_text_change_fails_closed():
    a, b = _docx_pair()
    b.tables[0].cell(0, 0).merge(b.tables[0].cell(0, 1))
    _fail_closed(a, b)


def test_docx_table_column_width_change_fails_closed():
    a, b = _docx_pair()
    b.tables[0].columns[0].width = 20_000_000
    _fail_closed(a, b)


def _png(rgb):
    # Two minimal valid PNGs; avoid a Pillow dependency in public CI.
    from base64 import b64decode
    a = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZQmcAAAAASUVORK5CYII="
    b = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    return b64decode(a if rgb == (20, 20, 20) else b)


def _image_pair():
    from docx import Document as Word
    import tempfile
    from pathlib import Path
    folder = Path(tempfile.mkdtemp(prefix="pywpdoc-image-safety-"))
    a, b = folder / "a.docx", folder / "b.docx"
    for path in (a, b):
        doc = Word()
        doc.add_paragraph("same body")
        doc.add_picture(BytesIO(_png((20, 20, 20))))
        doc.save(path)
    return Document.open(a), Document.open(b)


def test_docx_image_pixel_change_fails_closed():
    a, b = _image_pair()
    oldblob = a.backend.native.part.related_parts
    # A replacement at the same relId must be visible to the checker.
    media = [p for p in b.backend.native.part.related_parts.values()
             if p.content_type.startswith("image/")]
    assert len(media) == 1
    media[0]._blob = _png((200, 10, 30))
    _fail_closed(a, b)
    assert sha256(next(p.blob for p in oldblob.values() if p.content_type.startswith("image/"))).digest() != sha256(media[0].blob).digest()


def test_docx_image_resize_fails_closed():
    a, b = _image_pair()
    shape = b.backend.native.inline_shapes[0]
    shape.width = int(shape.width) + 200_000
    _fail_closed(a, b)


def test_docx_image_removed_fails_closed():
    a, b = _image_pair()
    drawing = b.backend.native.inline_shapes[0]._inline.getparent()
    drawing.getparent().remove(drawing)
    _fail_closed(a, b)


def test_docx_identical_table_and_identical_picture_are_allowed(tmp_path):
    a, b = _image_pair()
    a.add_table(1, 1).cell(0, 0).text = "same"
    b.add_table(1, 1).cell(0, 0).text = "same"
    # Only the selected body text changes.
    b.paragraphs[0].text = "changed body"
    out = a.compare(b, everywhere=False)
    target = tmp_path / "tracked_with_picture.docx"
    out.save(target)
    import zipfile
    with zipfile.ZipFile(target) as archive:
        assert b"<w:ins" in archive.read("word/document.xml")


def test_hwpx_table_border_reference_change_fails_closed():
    a, b = Document.new("hwpx"), Document.new("hwpx")
    for d in (a, b):
        d.add_paragraph("same body")
        d.add_table(1, 1).cell(0, 0).text = "same cell"
    b.tables[0].native._model.set("borderFillIDRef", "4")
    _fail_closed(a, b)


def test_hwpx_table_cell_span_change_fails_closed():
    a, b = Document.new("hwpx"), Document.new("hwpx")
    for d in (a, b):
        d.add_paragraph("same")
        d.add_table(1, 2).cell(0, 0).text = "same"
    b.tables[0].cell(0, 0).merge(b.tables[0].cell(0, 1))
    _fail_closed(a, b)


@pytest.mark.parametrize("fmt", ["docx", "hwpx"])
def test_unchanged_table_allows_body_text_revision(fmt, tmp_path):
    a, b = Document.new(fmt), Document.new(fmt)
    for d in (a, b):
        d.add_paragraph("old")
        d.add_table(1, 1).cell(0, 0).text = "same"
    b.paragraphs[0].text = "new"
    result = a.compare(b, everywhere=False)
    target = tmp_path / ("tracked." + fmt)
    result.save(target)
    import zipfile
    with zipfile.ZipFile(target) as archive:
        joined = b"".join(archive.read(n) for n in archive.namelist() if n.endswith(".xml"))
        assert (b"<w:ins" if fmt == "docx" else b"insertBegin") in joined


def test_hwpx_identical_picture_does_not_block_body_only_comparison():
    a, b = Document.new("hwpx"), Document.new("hwpx")
    for d in (a, b):
        d.add_paragraph("same")
        d.add_picture(_png((20, 20, 20)), "png")
    a.compare(b, everywhere=False)


def test_hwpx_picture_payload_change_fails_closed():
    a, b = Document.new("hwpx"), Document.new("hwpx")
    for d, pixels in ((a, (20, 20, 20)), (b, (200, 10, 30))):
        d.add_paragraph("same")
        d.add_picture(_png(pixels), "png")
    _fail_closed(a, b)


def test_hwpx_picture_resize_fails_closed():
    a, b = Document.new("hwpx"), Document.new("hwpx")
    for d in (a, b):
        d.add_paragraph("same")
        picture = d.add_picture(_png((20, 20, 20)), "png")
        if d is b:
            picture.width = 17000
    _fail_closed(a, b)
