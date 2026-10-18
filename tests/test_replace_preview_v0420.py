"""Read-only preview/parity and source-package immutability regressions."""
from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path

import pytest

from pywpdoc import Document, DocumentValueError
from pywpdoc.core.text import TextMap
from pywpdoc.features.replace import (
    ReplacementPreviewMatch, ReplacementPreviewReport, preview_replace,
)

ROOT = Path(__file__).resolve().parents[1]
BASELINE = json.loads((ROOT / 'tests/baseline/v0215_manifest.json').read_text(encoding='utf-8'))


def _part_hashes(path):
    with zipfile.ZipFile(path) as zf:
        return {name: hashlib.sha256(zf.read(name)).hexdigest() for name in zf.namelist()}


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_basic_counts_offsets_json_and_no_mutation(fmt):
    d = Document.new(fmt)
    p = d.add_paragraph('앞 needle needle 뒤')
    r = d.preview_replace('needle', '변경', everywhere=False)
    assert isinstance(r, ReplacementPreviewReport)
    assert r.has_matches and r.count == 2
    assert all(isinstance(m, ReplacementPreviewMatch) for m in r.matches)
    assert [(m.start, m.end) for m in r.matches] == [(2, 8), (9, 15)]
    assert [m.before for m in r.matches] == ['needle', 'needle']
    assert [m.after for m in r.matches] == ['변경', '변경']
    assert all(m.logical_before == '앞 needle needle 뒤' for m in r.matches)
    assert all(m.logical_after == '앞 변경 변경 뒤' for m in r.matches)
    assert p.text == '앞 needle needle 뒤'
    assert json.loads(r.to_json()) == r.to_dict()
    assert r.to_dict()['count'] == 2
    assert '변경' in r.to_json()
    assert d.replace_text('needle', '변경', everywhere=False) == r.count
    assert p.text == r.matches[0].logical_after


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_cross_run_match_and_preserved_format_is_not_changed_by_preview(fmt):
    d = Document.new(fmt)
    p = d.add_paragraph('')
    p.clear()
    a = p.add_run('prefix nee')
    a.bold = True
    b = p.add_run('dle end')
    b.italic = True
    r = d.preview_replace('needle', 'X', everywhere=False)
    assert [(m.start, m.end) for m in r.matches] == [(7, 13)]
    assert r.matches[0].logical_after == 'prefix X end'
    assert p.text == 'prefix needle end'
    assert a.bold is True and b.italic is True
    assert d.replace_text('needle', 'X', everywhere=False) == 1
    assert p.text == 'prefix X end'
    assert a.bold is True and b.italic is True


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_case_insensitive_limit_parity(fmt):
    d = Document.new(fmt)
    p = d.add_paragraph('Needle needle NEEDLE needle')
    r = d.preview_replace('needle', 'X', case_sensitive=False, count=2, everywhere=False)
    assert r.count == 2 and r.limit == 2
    assert [m.before for m in r.matches] == ['Needle', 'needle']
    assert r.matches[-1].logical_after == 'X X NEEDLE needle'
    assert p.text == 'Needle needle NEEDLE needle'
    assert d.replace_text('needle', 'X', case_sensitive=False, count=2, everywhere=False) == 2
    assert p.text == r.matches[0].logical_after


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
@pytest.mark.parametrize('count', [0, None, -1, -99, 1, 3])
def test_count_limits_match_replacement(fmt, count):
    d = Document.new(fmt)
    p = d.add_paragraph('aba aba aba')
    r = d.preview_replace('aba', '!', count=count, everywhere=False)
    assert r.limit == (None if count is None or count < 0 else count)
    assert r.count == (3 if count is None or count < 0 else min(3, count))
    if r.matches:
        final = r.matches[-1].logical_after
    else:
        final = p.text
    assert d.replace_text('aba', '!', count=count, everywhere=False) == r.count
    assert p.text == final


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_scope_and_traversal_order(fmt):
    d = Document.new(fmt)
    d.add_paragraph('body KEY')
    d.add_table(1, 1).cell(0, 0).text = 'table KEY'
    d.sections[0].header.text = 'header KEY'
    body = d.preview_replace('KEY', 'X', everywhere=False)
    full = d.preview_replace('KEY', 'X')
    assert body.count == 1
    assert full.count == 3
    assert [m.paragraph_index for m in full.matches] == sorted(m.paragraph_index for m in full.matches)
    assert all(m.logical_after.endswith('X') for m in full.matches)
    assert len(d.find('KEY')) == full.count


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_no_match_empty_result_is_serializable(fmt):
    d = Document.new(fmt)
    d.add_paragraph('abc')
    r = d.preview_replace('missing', 'X')
    assert not r.has_matches and r.count == 0
    assert r.to_dict()['matches'] == []
    assert r.to_json(indent=None).startswith('{')
    assert d.replace_text('missing', 'X') == 0


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_empty_search_same_error_code_as_replace(fmt):
    d = Document.new(fmt)
    with pytest.raises(DocumentValueError) as x:
        d.preview_replace('', 'X')
    with pytest.raises(DocumentValueError) as y:
        d.replace_text('', 'X')
    assert x.value.code == y.value.code == 'replace-query-empty'


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_strings_coerced_like_existing_replace(fmt):
    d = Document.new(fmt)
    p = d.add_paragraph('1 1')
    r = d.preview_replace(1, 2, everywhere=False)
    assert r.search == '1' and r.replacement == '2'
    assert r.count == 2
    assert r.matches[0].logical_after == '2 2'
    assert d.replace_text(1, 2, everywhere=False) == 2
    assert p.text == '2 2'


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_empty_replacement_and_overlapping_search(fmt):
    d = Document.new(fmt)
    p = d.add_paragraph('aaaa')
    r = d.preview_replace('aa', '', everywhere=False)
    assert r.count == 2
    assert [(m.start, m.end) for m in r.matches] == [(0, 2), (2, 4)]
    assert r.matches[0].logical_after == ''
    assert d.replace_text('aa', '', everywhere=False) == 2
    assert p.text == ''


def test_backend_direct_preview_and_module_independence():
    d = Document.new('hwpx')
    d.add_paragraph('backend KEY')
    r = preview_replace(d.backend, 'KEY', 'X', everywhere=False)
    assert r.count == 1 and r.matches[0].paragraph_index >= 0
    source = (ROOT/'src/pywpdoc/features/replace/preview.py').read_text(encoding='utf-8')
    for forbidden in ('pywpdoc.features.find', 'pywpdoc.hwpx', 'pywpdoc.docx', 'pywpdoc._engine'):
        assert forbidden not in source


def test_unsafe_docx_field_boundary_skipped_and_does_not_consume_limit():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    d = Document.new('docx')
    p = d.add_paragraph('')
    p.clear()
    p.add_run('nee')
    r = p.add_run('').native._r
    mark = OxmlElement('w:fldChar')
    mark.set(qn('w:fldCharType'), 'begin')
    r.append(mark)
    end_mark = OxmlElement('w:fldChar')
    end_mark.set(qn('w:fldCharType'), 'end')
    r.append(end_mark)
    p.add_run('dle')
    safe = d.add_paragraph('needle')
    assert TextMap.from_paragraph(p).unsafe_boundaries
    preview = d.preview_replace('needle', 'X', count=1, everywhere=False)
    assert preview.count == 1
    assert preview.matches[0].logical_before == 'needle'
    assert p.text == 'needle' and safe.text == 'needle'
    assert d.replace_text('needle', 'X', count=1, everywhere=False) == 1
    assert p.text == 'needle' and safe.text == 'X'


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_preview_then_actual_replacement_across_paragraphs(fmt):
    d = Document.new(fmt)
    ps = [d.add_paragraph(t) for t in ('needle needle', 'none', 'needle', 'needle')]
    report = d.preview_replace('needle', 'LONG', everywhere=False, count=3)
    assert report.count == 3
    assert [m.logical_before for m in report.matches] == ['needle needle', 'needle needle', 'needle']
    assert len({m.paragraph_index for m in report.matches}) == 2
    assert d.replace_text('needle', 'LONG', everywhere=False, count=3) == 3
    assert [p.text for p in ps] == ['LONG LONG', 'none', 'LONG', 'needle']


def test_original_hwpx_archive_is_unchanged_after_preview_and_resave(tmp_path):
    source = ROOT / BASELINE['source_fixture']
    before = _part_hashes(source)
    d = Document.open(source)
    assert d.preview_replace('쓰담', 'X').count
    out = tmp_path/'after.hwpx'
    d.save(out)
    assert _part_hashes(out) == before


def test_rebuilt_docx_archive_is_unchanged_after_preview_and_resave(tmp_path):
    source = ROOT / BASELINE['source_fixture']
    baseline = Document.rebuild(source, format='docx', preserve_layout=True)
    before_path = tmp_path/'before.docx'
    baseline.save(before_path)
    before = _part_hashes(before_path)
    d = Document.rebuild(source, format='docx', preserve_layout=True)
    assert d.preview_replace('쓰담', 'X').count
    after_path = tmp_path/'after.docx'
    d.save(after_path)
    assert _part_hashes(after_path) == before


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_preview_stays_read_only_after_save_and_reload(fmt, tmp_path):
    d = Document.new(fmt)
    d.add_paragraph('A A')
    path = tmp_path/('source.' + fmt)
    d.save(path)
    before = _part_hashes(path)
    reopened = Document.open(path)
    r = reopened.preview_replace('A', 'B', everywhere=False)
    assert r.count == 2
    out = tmp_path/('untouched.' + fmt)
    reopened.save(out)
    assert _part_hashes(out) == before
    assert Document.open(out).find('A', everywhere=False)


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
def test_report_snapshot_does_not_change_when_document_changes(fmt):
    d = Document.new(fmt)
    p = d.add_paragraph('A A')
    preview = d.preview_replace('A', 'B', everywhere=False)
    old = preview.to_dict()
    d.replace_text('A', 'B', everywhere=False)
    assert preview.to_dict() == old
    assert p.text == 'B B'


def test_json_control_characters_and_hangul_roundtrip():
    d = Document.new('docx')
    d.add_paragraph('가나\t가나')
    report = d.preview_replace('가나', '한글\n교체', everywhere=False)
    assert report.count == 2
    payload = json.loads(report.to_json())
    assert payload['matches'][0]['after'] == '한글\n교체'
    # TextMap's logical text deliberately omits native tab elements.
    assert payload['matches'][0]['logical_after'] == '한글\n교체한글\n교체'


def test_report_is_frozen_and_each_serialization_fresh():
    from dataclasses import FrozenInstanceError
    d = Document.new('docx')
    d.add_paragraph('abc')
    report = d.preview_replace('abc', 'def')
    with pytest.raises(FrozenInstanceError):
        report.search = 'no'
    with pytest.raises(FrozenInstanceError):
        report.matches[0].before = 'no'
    output = report.to_dict()
    output['matches'][0]['before'] = 'tampered'
    assert report.matches[0].before == 'abc'
