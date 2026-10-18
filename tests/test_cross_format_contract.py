"""Cross-format semantic contract baseline for the future unified API.

These tests intentionally use the native public APIs today (python-docx and
python-hwpx-next). When the common facade is introduced, the same semantic
cases should move behind that facade without changing the assertions.
"""
from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

import pytest
from lxml import etree


def _validate_docx(path: Path) -> None:
    from docx import Document as DocxDocument

    with ZipFile(path) as zf:
        assert zf.testzip() is None
        names = set(zf.namelist())
        assert "[Content_Types].xml" in names
        assert "word/document.xml" in names
        for name in names:
            if name.endswith((".xml", ".rels")):
                etree.fromstring(zf.read(name), etree.XMLParser(recover=False))
    DocxDocument(path)


def _save_and_reopen(kind: str, doc, path: Path):
    if kind == "hwpx":
        from pywpdoc._engine import Document, validate_hwpx_file

        doc.save(path)
        report = validate_hwpx_file(path)
        assert report.ok, report.issues
        return Document.open(path)
    doc.save(path)
    _validate_docx(path)
    from docx import Document as DocxDocument

    return DocxDocument(path)


@pytest.mark.parametrize("kind", ["hwpx", "docx"])
def test_contract_paragraph_run_mixed_text_roundtrip(kind: str, tmp_path: Path):
    if kind == "hwpx":
        from pywpdoc._engine import Document
        doc = Document.new()
    else:
        from docx import Document
        doc = Document()

    p = doc.add_paragraph()
    p.add_run("A\tB\nC")
    assert p.text == "A\tB\nC"

    reopened = _save_and_reopen(kind, doc, tmp_path / f"mixed.{kind}")
    assert reopened.paragraphs[-1].text == "A\tB\nC"


@pytest.mark.parametrize("kind", ["hwpx", "docx"])
def test_contract_insert_paragraph_before(kind: str, tmp_path: Path):
    if kind == "hwpx":
        from pywpdoc._engine import Document
        doc = Document.new()
    else:
        from docx import Document
        doc = Document()

    second = doc.add_paragraph("second")
    first = second.insert_paragraph_before("first")
    assert first.text == "first"

    reopened = _save_and_reopen(kind, doc, tmp_path / f"insert.{kind}")
    texts = [p.text for p in reopened.paragraphs]
    assert texts[-2:] == ["first", "second"]


@pytest.mark.parametrize("kind", ["hwpx", "docx"])
def test_contract_top_level_tables_exclude_nested_tables(kind: str, tmp_path: Path):
    if kind == "hwpx":
        from pywpdoc._engine import Document
        doc = Document.new()
        outer = doc.add_table(1, 1)
        outer.cell(0, 0).add_table(1, 1)
        assert len(doc.tables) == 1
        assert len(doc.all_tables) == 2
    else:
        from docx import Document
        doc = Document()
        outer = doc.add_table(1, 1)
        outer.cell(0, 0).add_table(1, 1)
        assert len(doc.tables) == 1

    reopened = _save_and_reopen(kind, doc, tmp_path / f"nested.{kind}")
    assert len(reopened.tables) == 1


@pytest.mark.parametrize("kind", ["hwpx", "docx"])
def test_contract_merged_row_cells_are_logical_grid_cells(kind: str, tmp_path: Path):
    if kind == "hwpx":
        from pywpdoc._engine import Document
        doc = Document.new()
        table = doc.add_table(2, 3)
        table.cell(0, 0).text = "merged"
        table.cell(0, 0).merge(table.cell(1, 1))
        assert len(table.rows[0].cells) == 3
        assert len(table.rows[1].cells) == 3
        assert table.rows[0].cells[0].element_model is table.rows[0].cells[1].element_model
        assert table.rows[1].cells[0].element_model is table.rows[0].cells[0].element_model
    else:
        from docx import Document
        doc = Document()
        table = doc.add_table(2, 3)
        table.cell(0, 0).text = "merged"
        table.cell(0, 0).merge(table.cell(1, 1))
        assert len(table.rows[0].cells) == 3
        assert len(table.rows[1].cells) == 3
        assert table.rows[0].cells[0]._tc is table.rows[0].cells[1]._tc
        assert table.rows[1].cells[0]._tc is table.rows[0].cells[0]._tc

    reopened = _save_and_reopen(kind, doc, tmp_path / f"merge.{kind}")
    assert len(reopened.tables[0].rows[0].cells) == 3


@pytest.mark.parametrize("kind", ["hwpx", "docx"])
def test_contract_core_properties(kind: str, tmp_path: Path):
    if kind == "hwpx":
        from pywpdoc._engine import Document
        doc = Document.new()
    else:
        from docx import Document
        doc = Document()

    doc.core_properties.title = "Unified Contract"
    doc.core_properties.author = "Tester"
    reopened = _save_and_reopen(kind, doc, tmp_path / f"props.{kind}")
    assert reopened.core_properties.title == "Unified Contract"
    assert reopened.core_properties.author == "Tester"


@pytest.mark.parametrize("kind", ["hwpx", "docx"])
def test_contract_section_margin_and_orientation(kind: str, tmp_path: Path):
    if kind == "hwpx":
        from pywpdoc._engine import Document, Mm, WD_ORIENT
        doc = Document.new()
        section = doc.sections[0]
        section.left_margin = Mm(18)
        section.orientation = WD_ORIENT.LANDSCAPE
        reopened = _save_and_reopen(kind, doc, tmp_path / "section.hwpx")
        assert abs(reopened.sections[0].left_margin.mm - 18.0) < 0.02
        assert reopened.sections[0].orientation == WD_ORIENT.LANDSCAPE
    else:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.shared import Mm
        doc = Document()
        section = doc.sections[0]
        section.left_margin = Mm(18)
        section.orientation = WD_ORIENT.LANDSCAPE
        reopened = _save_and_reopen(kind, doc, tmp_path / "section.docx")
        assert abs(reopened.sections[0].left_margin.mm - 18.0) < 0.02
        assert reopened.sections[0].orientation == WD_ORIENT.LANDSCAPE


@pytest.mark.parametrize("kind", ["hwpx", "docx"])
def test_contract_paragraph_style_next_style(kind: str, tmp_path: Path):
    if kind == "hwpx":
        from pywpdoc._engine import Document, WD_STYLE_TYPE
        doc = Document.new()
        a = doc.styles.add_style("Contract Question", WD_STYLE_TYPE.PARAGRAPH, clone_from="Body")
        b = doc.styles.add_style("Contract Answer", WD_STYLE_TYPE.PARAGRAPH, clone_from="Body")
    else:
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        doc = Document()
        a = doc.styles.add_style("Contract Question", WD_STYLE_TYPE.PARAGRAPH)
        b = doc.styles.add_style("Contract Answer", WD_STYLE_TYPE.PARAGRAPH)

    a.next_paragraph_style = b
    assert a.next_paragraph_style.name == "Contract Answer"
    reopened = _save_and_reopen(kind, doc, tmp_path / f"style.{kind}")
    assert reopened.styles["Contract Question"].next_paragraph_style.name == "Contract Answer"


def _docx_replace_paragraph_runs(paragraph, search: str, replacement: str) -> int:
    """Reference adapter for the future DocxBackend Paragraph.replace()."""
    if not search:
        raise ValueError("search must not be empty")
    runs = list(paragraph.runs)
    texts = [run.text for run in runs]
    combined = "".join(texts)
    starts = []
    pos = 0
    for text in texts:
        starts.append(pos)
        pos += len(text)
    matches = []
    cursor = 0
    while True:
        start = combined.find(search, cursor)
        if start < 0:
            break
        matches.append((start, start + len(search)))
        cursor = start + len(search)

    def locate(char_index: int):
        for i, (base, text) in enumerate(zip(starts, texts)):
            if base <= char_index < base + len(text):
                return i, char_index - base
        raise IndexError(char_index)

    for start, end in reversed(matches):
        first_i, first_off = locate(start)
        last_i, last_off = locate(end - 1)
        if first_i == last_i:
            text = runs[first_i].text
            runs[first_i].text = text[:first_off] + replacement + text[last_off + 1:]
        else:
            first_text = runs[first_i].text
            last_text = runs[last_i].text
            runs[first_i].text = first_text[:first_off] + replacement
            for i in range(first_i + 1, last_i):
                runs[i].text = ""
            runs[last_i].text = last_text[last_off + 1:]
    return len(matches)


@pytest.mark.parametrize("kind", ["hwpx", "docx"])
def test_contract_structure_preserving_replace_across_runs(kind: str, tmp_path: Path):
    if kind == "hwpx":
        from pywpdoc._engine import Document
        doc = Document.new()
        p = doc.add_paragraph()
        r1 = p.add_run("202")
        r1.bold = True
        r2 = p.add_run("6학년도")
        r2.italic = True
        assert p.replace("2026", "2027") == 1
    else:
        from docx import Document
        doc = Document()
        p = doc.add_paragraph()
        r1 = p.add_run("202")
        r1.bold = True
        r2 = p.add_run("6학년도")
        r2.italic = True
        assert _docx_replace_paragraph_runs(p, "2026", "2027") == 1

    assert p.text == "2027학년도"
    assert r1.bold is True
    assert r2.italic is True
    reopened = _save_and_reopen(kind, doc, tmp_path / f"replace.{kind}")
    rp = reopened.paragraphs[-1]
    visible = [run for run in rp.runs if run.text]
    assert rp.text == "2027학년도"
    assert visible[0].bold is True
    assert visible[1].italic is True


@pytest.mark.parametrize("kind", ["hwpx", "docx"])
def test_contract_paragraph_style_assignment_preserves_direct_run_formatting(kind: str, tmp_path: Path):
    if kind == "hwpx":
        from pywpdoc._engine import Document, WD_STYLE_TYPE, Pt
        doc = Document.new()
        a = doc.styles.add_style("Contract Style A", WD_STYLE_TYPE.PARAGRAPH, clone_from="Body")
        b = doc.styles.add_style("Contract Style B", WD_STYLE_TYPE.PARAGRAPH, clone_from="Body")
        a.font.italic = True
        b.font.size = Pt(15)
    else:
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        from docx.shared import Pt
        doc = Document()
        a = doc.styles.add_style("Contract Style A", WD_STYLE_TYPE.PARAGRAPH)
        b = doc.styles.add_style("Contract Style B", WD_STYLE_TYPE.PARAGRAPH)
        a.font.italic = True
        b.font.size = Pt(15)

    p = doc.add_paragraph("base", style=a)
    r = p.add_run(" direct")
    r.bold = True
    p.style = b
    assert r.bold is True
    assert p.style is not None and p.style.name == "Contract Style B"
    assert b.font.size is not None and abs(b.font.size.pt - 15.0) < 0.02
    p.style = None
    assert p.style is not None and p.style.name in {"Normal", "바탕글"}
    assert r.bold is True

    reopened = _save_and_reopen(kind, doc, tmp_path / f"style-assignment.{kind}")
    rp = reopened.paragraphs[-1]
    assert rp.style is not None and rp.style.name in {"Normal", "바탕글"}
    assert rp.runs[-1].bold is True


@pytest.mark.parametrize("kind", ["hwpx", "docx"])
def test_contract_clear_preserves_paragraph_style_and_removes_visible_content(kind: str, tmp_path: Path):
    if kind == "hwpx":
        from pywpdoc._engine import Document, WD_STYLE_TYPE
        doc = Document.new()
        style = doc.styles.add_style("Contract Clear", WD_STYLE_TYPE.PARAGRAPH, clone_from="Body")
    else:
        from docx import Document
        from docx.enum.style import WD_STYLE_TYPE
        doc = Document()
        style = doc.styles.add_style("Contract Clear", WD_STYLE_TYPE.PARAGRAPH)

    p = doc.add_paragraph("before", style=style)
    returned = p.clear()
    assert returned is p
    assert p.text == ""
    assert p.style is not None and p.style.name == "Contract Clear"
    p.add_run("after")

    reopened = _save_and_reopen(kind, doc, tmp_path / f"clear.{kind}")
    rp = reopened.paragraphs[-1]
    assert rp.text == "after"
    assert rp.style is not None and rp.style.name == "Contract Clear"
