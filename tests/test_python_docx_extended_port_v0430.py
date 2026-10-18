"""Behavioral ports of selected python-docx publicly observable tests.

Upstream reference categories: tests/test_document.py, tests/test_table.py,
features/steps/table.py and public Quickstart. These are *adaptations*, not a
verbatim copy or claim of complete upstream test compatibility.
"""
from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile
import pytest
from docx import Document as RefDocument
from pywpdoc import Document


def _pair(fmt, tmp_path):
    return RefDocument(), Document.new(fmt)


def _roundtrip(fmt, tmp_path, ref, candidate):
    a, b = tmp_path/'reference.docx', tmp_path/f'candidate.{fmt}'
    ref.save(a); candidate.save(b)
    with ZipFile(b) as z: assert z.testzip() is None
    return RefDocument(a), Document.open(b)


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
@pytest.mark.parametrize('content', ['기본 문단', 'English with spaces', 'A\tB\nC', '혼합 2026-10-14 Latin'])
def test_upstream_document_add_paragraph_roundtrip(fmt, content, tmp_path):
    r, c = _pair(fmt, tmp_path)
    r.add_paragraph(content); c.add_paragraph(content)
    rr, cc = _roundtrip(fmt, tmp_path, r, c)
    assert [p.text for p in rr.paragraphs] == [p.text for p in cc.paragraphs]


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
@pytest.mark.parametrize('sequence', [('A','B'), ('앞','중','뒤'), ('01','-','2026','-','10')])
def test_upstream_paragraph_runs_concatenate(fmt, sequence, tmp_path):
    r, c = _pair(fmt, tmp_path)
    for d in (r, c):
        p=d.add_paragraph()
        for text in sequence: p.add_run(text)
    rr, cc = _roundtrip(fmt, tmp_path, r, c)
    assert rr.paragraphs[-1].text == cc.paragraphs[-1].text == ''.join(sequence)


@pytest.mark.parametrize('fmt', ['docx', 'hwpx'])
@pytest.mark.parametrize('style', [(True,False),(False,True),(True,True)])
def test_upstream_run_bold_italic_semantics(fmt, style, tmp_path):
    r, c = _pair(fmt,tmp_path)
    for d in (r,c):
        run=d.add_paragraph().add_run('강조 본문')
        run.bold,run.italic=style
    rr,cc=_roundtrip(fmt,tmp_path,r,c)
    def visible(d):
        return [(x.text,bool(x.bold),bool(x.italic)) for x in d.paragraphs[-1].runs if x.text]
    assert visible(rr)==visible(cc)==[('강조 본문',*style)]


@pytest.mark.parametrize('fmt', ['docx','hwpx'])
@pytest.mark.parametrize('replacement', ['갱신','A\tB','두 줄\n본문'])
def test_upstream_paragraph_clear_rebuild(fmt,replacement,tmp_path):
    r,c=_pair(fmt,tmp_path)
    for d in (r,c):
        p=d.add_paragraph()
        p.add_run('old');p.add_run(' and gone');p.clear();p.add_run(replacement)
    rr,cc=_roundtrip(fmt,tmp_path,r,c)
    assert rr.paragraphs[-1].text==cc.paragraphs[-1].text==replacement


@pytest.mark.parametrize('fmt',['docx','hwpx'])
@pytest.mark.parametrize('rows,cols', [(1,1),(2,3),(4,2)])
def test_upstream_table_grid_access(fmt,rows,cols,tmp_path):
    r,c=_pair(fmt,tmp_path)
    for d in (r,c):
        t=d.add_table(rows,cols)
        for ri in range(rows):
            for ci in range(cols): t.cell(ri,ci).text=f'R{ri}C{ci}'
    rr,cc=_roundtrip(fmt,tmp_path,r,c)
    def grid(doc):
        return [[cell.text for cell in row.cells] for row in doc.tables[0].rows]
    assert grid(rr)==grid(cc)


@pytest.mark.parametrize('fmt',['docx','hwpx'])
@pytest.mark.parametrize('extra', [1,2,3])
def test_upstream_add_table_rows(fmt,extra,tmp_path):
    r,c=_pair(fmt,tmp_path)
    for d in (r,c):
        t=d.add_table(1,2)
        t.cell(0,0).text='항목';t.cell(0,1).text='값'
        for k in range(extra):
            row=t.add_row(); row.cells[0].text=f'K{k}';row.cells[1].text=f'V{k}'
    rr,cc=_roundtrip(fmt,tmp_path,r,c)
    def grid(doc):return [[x.text for x in row.cells] for row in doc.tables[0].rows]
    assert grid(rr)==grid(cc)


@pytest.mark.parametrize('fmt',['docx','hwpx'])
@pytest.mark.parametrize('axis', ['horizontal','vertical'])
def test_upstream_merge_cell_semantics(fmt,axis,tmp_path):
    r,c=_pair(fmt,tmp_path)
    for d in (r,c):
        t=d.add_table(2,2);t.cell(0,0).text='병합'
        end=(0,1) if axis=='horizontal' else (1,0)
        t.cell(0,0).merge(t.cell(*end))
    rr,cc=_roundtrip(fmt,tmp_path,r,c)
    def grid(doc):return [[x.text for x in row.cells] for row in doc.tables[0].rows]
    assert grid(rr)==grid(cc)


@pytest.mark.parametrize('fmt',['docx','hwpx'])
def test_upstream_nested_table_access(fmt,tmp_path):
    r,c=_pair(fmt,tmp_path)
    for d in (r,c):
        outer=d.add_table(1,1); outer.cell(0,0).text='Outer'
        outer.cell(0,0).add_table(2,1).cell(1,0).text='Inner target'
    rr,cc=_roundtrip(fmt,tmp_path,r,c)
    assert rr.tables[0].cell(0,0).tables[0].cell(1,0).text=='Inner target'
    assert cc.tables[0].cell(0,0).tables[0].cell(1,0).text=='Inner target'


@pytest.mark.parametrize('fmt',['docx','hwpx'])
@pytest.mark.parametrize('field,value',[('title','확장 QA'),('subject','Common API 비교'),('author','QA Operator')])
def test_upstream_metadata_roundtrip(fmt,field,value,tmp_path):
    r,c=_pair(fmt,tmp_path)
    for d in (r,c): setattr(d.core_properties,field,value)
    rr,cc=_roundtrip(fmt,tmp_path,r,c)
    assert getattr(rr.core_properties,field)==getattr(cc.core_properties,field)==value


@pytest.mark.parametrize('fmt',['docx','hwpx'])
def test_upstream_insert_paragraph_before_sequence(fmt,tmp_path):
    r,c=_pair(fmt,tmp_path)
    for d in (r,c):
        d.add_paragraph('세 번째').insert_paragraph_before('두 번째').insert_paragraph_before('첫 번째')
    rr,cc=_roundtrip(fmt,tmp_path,r,c)
    assert [x.text for x in rr.paragraphs]==[x.text for x in cc.paragraphs]
